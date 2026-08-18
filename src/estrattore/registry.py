from __future__ import annotations

import json
import re
import shutil
from dataclasses import asdict, dataclass, replace
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from estrattore.fiscal_validation import (
    normalize_italian_tax_code,
    normalize_italian_vat_number,
)
from estrattore.ocr_normalization import normalize_company_legal_form
from estrattore.runtime_paths import get_runtime_paths


REGISTRY_SCHEMA_VERSION = 1
REGISTRY_EXPECTED_FILE_DESCRIPTION = (
    "L'anagrafica non e nel formato aspettato."
)
APP_REGISTRY_EXCEL_HEADERS = [
    "codice_anagrafica",
    "nominativo",
    "codice_fiscale",
    "partita_iva",
    "tipologia",
    "ultima_modifica",
    "utente",
]
_RUNTIME_PATHS = get_runtime_paths()
DEFAULT_REGISTRY_INPUT_PATH = _RUNTIME_PATHS.registry_input_file
DEFAULT_REGISTRY_CACHE_PATH = _RUNTIME_PATHS.registry_cache_file
DEFAULT_REGISTRY_DUPLICATE_REPORT_PATH = (
    _RUNTIME_PATHS.registry_duplicate_report_json_file
)


@dataclass(frozen=True)
class RegistryRecord:
    registry_id: str
    source_row: int
    codice_anagrafica: str | None
    nominativo: str
    normalized_name: str
    codice_fiscale_raw: str
    codice_fiscale: str | None
    partita_iva_raw: str
    partita_iva: str | None
    tipologia: str
    ultima_modifica: str | None = None
    utente: str | None = None
    warnings: list[str] | None = None


@dataclass(frozen=True)
class RegistryCache:
    schema_version: int
    source_file: str
    records: list[RegistryRecord]


@dataclass(frozen=True)
class RegistryCleaningResult:
    cache: RegistryCache
    report: dict[str, Any]


def build_registry_cache_from_excel(
    excel_path: Path = DEFAULT_REGISTRY_INPUT_PATH,
    output_path: Path | None = DEFAULT_REGISTRY_CACHE_PATH,
) -> RegistryCache:
    records = read_registry_records_from_excel(excel_path)
    cache = RegistryCache(
        schema_version=REGISTRY_SCHEMA_VERSION,
        source_file=excel_path.as_posix(),
        records=records,
    )

    if output_path is not None:
        save_registry_cache(cache, output_path)

    return cache


def import_registry_excel(
    source_path: Path,
    input_path: Path = DEFAULT_REGISTRY_INPUT_PATH,
    cache_path: Path = DEFAULT_REGISTRY_CACHE_PATH,
) -> RegistryCache:
    if source_path.suffix.lower() != ".xlsx":
        raise ValueError("Il file anagrafica deve essere un Excel .xlsx")

    input_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.parent.mkdir(parents=True, exist_ok=True)

    temporary_input_path = input_path.with_name(f".{input_path.stem}.importing.xlsx")
    temporary_cache_path = cache_path.with_name(f".{cache_path.stem}.importing.json")
    try:
        shutil.copy2(source_path, temporary_input_path)
        records = read_registry_records_from_excel(temporary_input_path)
        cache = RegistryCache(
            schema_version=REGISTRY_SCHEMA_VERSION,
            source_file=input_path.as_posix(),
            records=records,
        )
        save_registry_cache(cache, temporary_cache_path)

        temporary_input_path.replace(input_path)
        temporary_cache_path.replace(cache_path)
        DEFAULT_REGISTRY_DUPLICATE_REPORT_PATH.unlink(missing_ok=True)
        return cache
    finally:
        temporary_input_path.unlink(missing_ok=True)
        temporary_cache_path.unlink(missing_ok=True)


def read_registry_records_from_excel(excel_path: Path) -> list[RegistryRecord]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        if _is_app_registry_worksheet(worksheet):
            return _read_app_registry_records_from_worksheet(worksheet)

        _validate_registry_worksheet_columns(worksheet)
        records = []

        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_index <= 11:
                continue

            values = list(row[:13])
            values.extend([None] * (13 - len(values)))
            if not any(_cell_text(value) for value in values):
                continue

            nominativo = _cell_text(values[0])
            codice_anagrafica = extract_registry_code(nominativo)
            extracted_codice_fiscale, extracted_partita_iva = (
                extract_registry_fiscal_identifiers(nominativo)
            )
            normalized_name = normalize_registry_name(nominativo)
            codice_fiscale_raw = _cell_text(values[4])
            partita_iva_raw = _cell_text(values[5])
            codice_fiscale = (
                _normalized_registry_tax_code(codice_fiscale_raw)
                or extracted_codice_fiscale
            )
            partita_iva = (
                _normalized_registry_vat_number(partita_iva_raw)
                or extracted_partita_iva
            )

            records.append(
                RegistryRecord(
                    registry_id=_registry_id_from_nominativo(nominativo, row_index),
                    source_row=row_index,
                    codice_anagrafica=codice_anagrafica,
                    nominativo=nominativo,
                    normalized_name=normalized_name,
                    codice_fiscale_raw=codice_fiscale_raw,
                    codice_fiscale=codice_fiscale,
                    partita_iva_raw=partita_iva_raw,
                    partita_iva=partita_iva,
                    tipologia=_cell_text(values[6]),
                    ultima_modifica=_json_text(values[7]),
                    utente=_cell_text(values[12]) or None,
                    warnings=_registry_record_warnings(
                        normalized_name=normalized_name,
                        codice_fiscale_raw=codice_fiscale_raw,
                        codice_fiscale=codice_fiscale,
                        partita_iva_raw=partita_iva_raw,
                        partita_iva=partita_iva,
                    ),
                )
            )

        return records
    finally:
        workbook.close()


def clean_registry_duplicate_records(
    records: list[RegistryRecord],
    duplicate_report_path: Path = DEFAULT_REGISTRY_DUPLICATE_REPORT_PATH,
    selected_merge_group_keys: set[str] | None = None,
) -> list[RegistryRecord]:
    result = clean_registry_duplicate_records_with_report(
        records,
        selected_merge_group_keys=selected_merge_group_keys,
    )
    if result.report["merged_count"] or result.report["conflict_count"]:
        _save_registry_duplicate_report(result.report, duplicate_report_path)
    else:
        duplicate_report_path.unlink(missing_ok=True)

    return result.cache.records


def clean_registry_cache_duplicates(
    cache: RegistryCache,
    duplicate_report_path: Path = DEFAULT_REGISTRY_DUPLICATE_REPORT_PATH,
    selected_merge_group_keys: set[str] | None = None,
) -> RegistryCleaningResult:
    result = clean_registry_duplicate_records_with_report(
        cache.records,
        selected_merge_group_keys=selected_merge_group_keys,
    )
    cleaned_cache = RegistryCache(
        schema_version=cache.schema_version,
        source_file=cache.source_file,
        records=result.cache.records,
    )

    if result.report["merged_count"] or result.report["conflict_count"]:
        _save_registry_duplicate_report(result.report, duplicate_report_path)
    else:
        duplicate_report_path.unlink(missing_ok=True)

    return RegistryCleaningResult(cache=cleaned_cache, report=result.report)


def preview_registry_cache_duplicates(cache: RegistryCache) -> RegistryCleaningResult:
    result = clean_registry_duplicate_records_with_report(cache.records)
    preview_cache = RegistryCache(
        schema_version=cache.schema_version,
        source_file=cache.source_file,
        records=result.cache.records,
    )
    return RegistryCleaningResult(cache=preview_cache, report=result.report)


def clean_registry_duplicate_records_with_report(
    records: list[RegistryRecord],
    selected_merge_group_keys: set[str] | None = None,
) -> RegistryCleaningResult:
    duplicate_groups = _registry_duplicate_groups(records)
    duplicate_report: list[dict[str, Any]] = []
    cleaned_records: list[RegistryRecord] = []
    consumed_indexes: set[int] = set()

    for record_index, record in enumerate(records):
        if record_index in consumed_indexes:
            continue

        group = next(
            (candidate for candidate in duplicate_groups if record_index in candidate),
            None,
        )
        if group is None:
            cleaned_records.append(record)
            continue

        consumed_indexes.update(group)
        group_records = [records[index] for index in sorted(group)]
        group_key = _registry_duplicate_group_key(group_records)
        merge_is_selected = (
            selected_merge_group_keys is None
            or group_key in selected_merge_group_keys
        )
        if _registry_records_are_mergeable(group_records) and merge_is_selected:
            merged_record = _merge_registry_duplicate_records(group_records)
            cleaned_records.append(
                _record_with_warning(
                    merged_record,
                    "merged_duplicate_fiscal_identifier",
                )
            )
            duplicate_report.append(
                _registry_duplicate_report_item("merged", group_records, [merged_record])
            )
            continue

        conflict_records = [
            _record_with_warning(record, "conflicting_duplicate_fiscal_identifier")
            for record in group_records
        ]
        cleaned_records.extend(conflict_records)
        duplicate_report.append(
            _registry_duplicate_report_item("conflict", group_records, conflict_records)
        )

    report = _registry_duplicate_report_payload(duplicate_report)
    cache = RegistryCache(
        schema_version=REGISTRY_SCHEMA_VERSION,
        source_file="",
        records=cleaned_records,
    )

    return RegistryCleaningResult(cache=cache, report=report)


def _registry_duplicate_groups(records: list[RegistryRecord]) -> list[set[int]]:
    indexes_by_fiscal_identifier: dict[str, list[int]] = {}

    for record_index, record in enumerate(records):
        for fiscal_identifier in _record_fiscal_identifiers(record):
            indexes_by_fiscal_identifier.setdefault(fiscal_identifier, []).append(record_index)

    parent = list(range(len(records)))

    def find(index: int) -> int:
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    def union(left: int, right: int) -> None:
        left_root = find(left)
        right_root = find(right)
        if left_root != right_root:
            parent[right_root] = left_root

    for indexes in indexes_by_fiscal_identifier.values():
        if len(indexes) < 2:
            continue
        first_index = indexes[0]
        for index in indexes[1:]:
            union(first_index, index)

    groups_by_root: dict[int, set[int]] = {}
    duplicate_indexes = {
        index
        for indexes in indexes_by_fiscal_identifier.values()
        if len(indexes) > 1
        for index in indexes
    }
    for index in duplicate_indexes:
        groups_by_root.setdefault(find(index), set()).add(index)

    return [group for group in groups_by_root.values() if len(group) > 1]


def _record_fiscal_identifiers(record: RegistryRecord) -> list[str]:
    identifiers = []
    if record.codice_fiscale:
        identifiers.append(f"cf:{record.codice_fiscale}")
    if record.partita_iva:
        identifiers.append(f"piva:{record.partita_iva}")
    return identifiers


def _registry_records_are_mergeable(records: list[RegistryRecord]) -> bool:
    if _records_are_same_except_registry_code_and_missing_vat(records):
        return True

    return (
        _record_registry_codes_are_compatible(records)
        and _record_names_are_compatible(records)
        and all(
            _record_field_values_are_compatible(values)
            for values in (
                [record.codice_fiscale for record in records],
                [record.partita_iva for record in records],
                [record.tipologia for record in records],
            )
        )
    )


def _record_registry_codes_are_compatible(records: list[RegistryRecord]) -> bool:
    registry_codes = {
        record.codice_anagrafica
        for record in records
        if record.codice_anagrafica
    }
    return len(registry_codes) <= 1


def _records_are_same_except_registry_code_and_missing_vat(
    records: list[RegistryRecord],
) -> bool:
    if len(records) != 2:
        return False

    left, right = records
    if not left.codice_anagrafica or not right.codice_anagrafica:
        return False
    if left.codice_anagrafica == right.codice_anagrafica:
        return False

    if not _record_names_are_compatible(records):
        return False
    if not _record_field_values_are_compatible([left.codice_fiscale, right.codice_fiscale]):
        return False
    if not _record_field_values_are_compatible([left.tipologia, right.tipologia]):
        return False

    vat_values = [left.partita_iva, right.partita_iva]
    return sum(1 for value in vat_values if value) == 1


def _record_names_are_compatible(records: list[RegistryRecord]) -> bool:
    registry_codes = {record.codice_anagrafica for record in records if record.codice_anagrafica}
    fiscal_identifiers = {
        identifier.split(":", 1)[1]
        for record in records
        for identifier in _record_fiscal_identifiers(record)
    }
    names = [
        record.normalized_name
        for record in records
        if record.normalized_name
        and record.normalized_name not in registry_codes
        and record.normalized_name not in fiscal_identifiers
    ]
    return _record_field_values_are_compatible(names)


def _record_field_values_are_compatible(values: list[str | None]) -> bool:
    normalized_values = [_normalize_registry_merge_value(value) for value in values]
    unique_values = sorted({value for value in normalized_values if value})
    if len(unique_values) <= 1:
        return True

    longest_value = max(unique_values, key=len)
    return all(_registry_value_is_contained_variant(value, longest_value) for value in unique_values)


def _normalize_registry_merge_value(value: str | None) -> str:
    return " ".join((value or "").upper().split())


def _registry_value_is_contained_variant(shorter_value: str, longer_value: str) -> bool:
    if shorter_value == longer_value:
        return True

    if len(shorter_value) < 8:
        return False

    return shorter_value in longer_value


def _merge_registry_duplicate_records(records: list[RegistryRecord]) -> RegistryRecord:
    selected_record = max(
        records,
        key=lambda record: (
            _record_completeness_score(record),
            -record.source_row,
        ),
    )

    return replace(
        selected_record,
        codice_anagrafica=selected_record.codice_anagrafica,
        nominativo=_best_record_text_value(records, "nominativo") or selected_record.nominativo,
        normalized_name=(
            _best_registry_name_value(records)
            or selected_record.normalized_name
        ),
        codice_fiscale_raw=(
            _best_record_text_value(records, "codice_fiscale_raw")
            or selected_record.codice_fiscale_raw
        ),
        codice_fiscale=_best_record_text_value(records, "codice_fiscale"),
        partita_iva_raw=(
            _best_record_text_value(records, "partita_iva_raw")
            or selected_record.partita_iva_raw
        ),
        partita_iva=_best_record_text_value(records, "partita_iva"),
        tipologia=_best_record_text_value(records, "tipologia") or selected_record.tipologia,
        ultima_modifica=(
            _best_record_text_value(records, "ultima_modifica")
            or selected_record.ultima_modifica
        ),
        utente=_best_record_text_value(records, "utente") or selected_record.utente,
        warnings=_merged_record_warnings(records),
    )


def _record_completeness_score(record: RegistryRecord) -> int:
    return _record_substantive_identity_score(record) * 10 + sum(
        1
        for value in (
            record.ultima_modifica,
            record.utente,
        )
        if value
    )


def _record_substantive_identity_score(
    record: RegistryRecord,
    group_records: list[RegistryRecord] | None = None,
) -> int:
    score = 0
    if _record_has_substantive_name(record, group_records or [record]):
        score += 1
    if record.tipologia:
        score += 1
    score += len(_record_distinct_fiscal_identifier_values(record))
    return score


def _record_distinct_fiscal_identifier_values(record: RegistryRecord) -> set[str]:
    return {
        value
        for value in (record.codice_fiscale, record.partita_iva)
        if value
    }


def _record_has_substantive_name(
    record: RegistryRecord,
    group_records: list[RegistryRecord],
) -> bool:
    if not record.normalized_name:
        return False

    registry_codes = {
        candidate.codice_anagrafica
        for candidate in group_records
        if candidate.codice_anagrafica
    }
    fiscal_identifiers = {
        value
        for candidate in group_records
        for value in _record_distinct_fiscal_identifier_values(candidate)
    }
    return (
        record.normalized_name not in registry_codes
        and record.normalized_name not in fiscal_identifiers
    )


def _best_record_text_value(records: list[RegistryRecord], field_name: str) -> str | None:
    values = [
        str(value).strip()
        for record in records
        if (value := getattr(record, field_name)) is not None and str(value).strip()
    ]
    if not values:
        return None

    return max(values, key=lambda value: (len(value), -values.index(value)))


def _best_registry_name_value(records: list[RegistryRecord]) -> str | None:
    registry_codes = {record.codice_anagrafica for record in records if record.codice_anagrafica}
    fiscal_identifiers = {
        identifier.split(":", 1)[1]
        for record in records
        for identifier in _record_fiscal_identifiers(record)
    }
    values = [
        record.normalized_name
        for record in records
        if record.normalized_name
        and record.normalized_name not in registry_codes
        and record.normalized_name not in fiscal_identifiers
    ]
    if not values:
        return _best_record_text_value(records, "normalized_name")

    return max(values, key=lambda value: (len(value), -values.index(value)))


def _merged_record_warnings(records: list[RegistryRecord]) -> list[str]:
    warnings = []
    for record in records:
        for warning in record.warnings or []:
            if warning not in warnings:
                warnings.append(warning)
    return warnings


def _record_with_warning(record: RegistryRecord, warning: str) -> RegistryRecord:
    warnings = list(record.warnings or [])
    if warning not in warnings:
        warnings.append(warning)
    return replace(record, warnings=warnings)


def _registry_duplicate_report_item(
    status: str,
    source_records: list[RegistryRecord],
    output_records: list[RegistryRecord],
) -> dict[str, Any]:
    fiscal_identifiers = sorted(
        {
            identifier
            for record in source_records
            for identifier in _record_fiscal_identifiers(record)
        }
    )
    return {
        "group_key": _registry_duplicate_group_key(source_records),
        "status": status,
        "fiscal_identifiers": fiscal_identifiers,
        "source_rows": [record.source_row for record in source_records],
        "source_records": [_registry_duplicate_report_record(record) for record in source_records],
        "output_records": [_registry_duplicate_report_record(record) for record in output_records],
        "kept_records": [_registry_duplicate_report_record(record) for record in output_records],
        "deleted_records": [
            _registry_duplicate_report_record(record)
            for record in source_records
            if status == "merged"
            and not any(
                record.registry_id == output_record.registry_id
                and record.source_row == output_record.source_row
                for output_record in output_records
            )
        ],
    }


def _registry_duplicate_group_key(records: list[RegistryRecord]) -> str:
    return "|".join(str(record.source_row) for record in sorted(records, key=lambda record: record.source_row))


def _registry_duplicate_report_record(record: RegistryRecord) -> dict[str, Any]:
    return {
        "registry_id": record.registry_id,
        "source_row": record.source_row,
        "codice_anagrafica": record.codice_anagrafica,
        "nominativo": record.nominativo,
        "normalized_name": record.normalized_name,
        "codice_fiscale_raw": record.codice_fiscale_raw,
        "codice_fiscale": record.codice_fiscale,
        "partita_iva_raw": record.partita_iva_raw,
        "partita_iva": record.partita_iva,
        "tipologia": record.tipologia,
        "warnings": record.warnings or [],
    }


def _save_registry_duplicate_report(
    duplicate_report: dict[str, Any],
    duplicate_report_path: Path,
) -> None:
    duplicate_report_path.parent.mkdir(parents=True, exist_ok=True)
    duplicate_report_path.write_text(
        json.dumps(duplicate_report, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def _registry_duplicate_report_payload(
    duplicate_report: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "merged_count": sum(1 for item in duplicate_report if item["status"] == "merged"),
        "conflict_count": sum(1 for item in duplicate_report if item["status"] == "conflict"),
        "deleted_count": sum(len(item["deleted_records"]) for item in duplicate_report),
        "groups": duplicate_report,
    }


def save_registry_duplicate_report_excel(
    duplicate_report: dict[str, Any],
    output_path: Path,
) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Pulizia anagrafica"
    grouped_report = {
        "Righe unite automaticamente": [
            group
            for group in duplicate_report.get("groups", [])
            if group.get("status") == "merged"
        ],
        "Duplicati da controllare": [
            group
            for group in duplicate_report.get("groups", [])
            if group.get("status") == "conflict"
        ],
    }

    for section_title, groups in grouped_report.items():
        if worksheet.max_row > 1 or worksheet.cell(row=1, column=1).value:
            worksheet.append([])

        worksheet.append([section_title])
        section_row = worksheet.max_row
        worksheet.cell(row=section_row, column=1).font = Font(bold=True, size=13)
        worksheet.cell(row=section_row, column=1).fill = PatternFill(
            fill_type="solid",
            fgColor="DDEFE5",
        )
        worksheet.append(
            [
                "Gruppo",
                "Ruolo",
                "Riga Excel",
                "Codice anagrafica",
                "Nominativo",
                "Nome normalizzato",
                "Codice fiscale",
                "Partita IVA",
                "Tipologia",
                "Identificativi duplicati",
            ]
        )
        header_row = worksheet.max_row
        for cell in worksheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fill_type="solid", fgColor="EEF3EF")

        for group_index, group in enumerate(groups, start=1):
            fiscal_identifiers = ", ".join(group.get("fiscal_identifiers") or [])
            if group.get("status") == "merged":
                rows = [
                    ("eliminata", record)
                    for record in group.get("deleted_records", [])
                ] + [
                    ("tenuta", record)
                    for record in group.get("kept_records", [])
                ]
            else:
                rows = [
                    ("da controllare", record)
                    for record in group.get("source_records", [])
                ]

            for role, record in rows:
                worksheet.append(
                    [
                        group_index,
                        role,
                        record.get("source_row"),
                        record.get("codice_anagrafica") or "",
                        record.get("nominativo") or "",
                        record.get("normalized_name") or "",
                        record.get("codice_fiscale") or "",
                        record.get("partita_iva") or "",
                        record.get("tipologia") or "",
                        fiscal_identifiers,
                    ]
                )

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length + 2, 12),
            48,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path


def save_registry_cache_excel(cache: RegistryCache, output_path: Path) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Anagrafica estrattore scan"
    worksheet.append(APP_REGISTRY_EXCEL_HEADERS)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="DDEFE5")

    for record in cache.records:
        worksheet.append(
            [
                record.codice_anagrafica or "",
                record.nominativo,
                record.codice_fiscale or "",
                record.partita_iva or "",
                record.tipologia,
                record.ultima_modifica or "",
                record.utente or "",
            ]
        )

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length + 2, 12),
            48,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path


def save_registry_cache(cache: RegistryCache, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(_cache_to_jsonable(cache), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return output_path


def load_registry_cache(cache_path: Path = DEFAULT_REGISTRY_CACHE_PATH) -> RegistryCache:
    data = json.loads(cache_path.read_text(encoding="utf-8"))
    return RegistryCache(
        schema_version=int(data["schema_version"]),
        source_file=str(data["source_file"]),
        records=[
            RegistryRecord(
                registry_id=str(record["registry_id"]),
                source_row=int(record["source_row"]),
                codice_anagrafica=record.get("codice_anagrafica")
                or extract_registry_code(str(record["nominativo"])),
                nominativo=str(record["nominativo"]),
                normalized_name=str(record["normalized_name"]),
                codice_fiscale_raw=str(record.get("codice_fiscale_raw", "")),
                codice_fiscale=record.get("codice_fiscale"),
                partita_iva_raw=str(record.get("partita_iva_raw", "")),
                partita_iva=record.get("partita_iva"),
                tipologia=str(record["tipologia"]),
                ultima_modifica=record.get("ultima_modifica"),
                utente=record.get("utente"),
                warnings=record.get("warnings") or [],
            )
            for record in data.get("records", [])
        ],
    )


def normalize_registry_name(value: str) -> str:
    text = _cell_text(value)
    text = re.sub(r"^\d{3,}(?:\s+|$)", "", text)
    text = _remove_registry_fiscal_identifiers_from_text(text)
    text = re.sub(
        r"\b(?:P\.?\s*IVA|PARTITA\s+IVA|C\.?\s*F\.?|COD(?:ICE)?\.?\s+FISCALE)\b",
        " ",
        text,
        flags=re.IGNORECASE,
    )
    text = text.replace('"', " ")
    text = text.replace(";", " ")
    text = re.sub(r"[*#]+", " ", text)
    text = normalize_company_legal_form(text)
    text = re.sub(r"[^0-9A-Za-zÀ-ÖØ-öø-ÿ&.' /-]+", " ", text)
    return " ".join(text.upper().split())


def extract_registry_code(value: str) -> str | None:
    text = _cell_text(value)
    match = re.match(r"^(\d{3,})(?:\b|\s|$)", text)
    if not match:
        return None

    return match.group(1)


def extract_registry_fiscal_identifiers(value: str) -> tuple[str | None, str | None]:
    codice_fiscale: str | None = None
    partita_iva: str | None = None

    for candidate in _registry_fiscal_identifier_candidates(value):
        codice_fiscale = codice_fiscale or _normalized_registry_tax_code(candidate)
        partita_iva = partita_iva or _normalized_registry_vat_number(candidate)
        if codice_fiscale and partita_iva:
            break

    return codice_fiscale, partita_iva


def _remove_registry_fiscal_identifiers_from_text(value: str) -> str:
    text = value
    for candidate in _registry_fiscal_identifier_candidates(value):
        if (
            _normalized_registry_tax_code(candidate) is None
            and _normalized_registry_vat_number(candidate) is None
        ):
            continue

        text = text.replace(candidate, " ")

    return text


def _registry_fiscal_identifier_candidates(value: str) -> list[str]:
    return re.findall(r"\b(?:IT)?[A-Z0-9]{11,16}\b", _cell_text(value).upper())


def _validate_registry_worksheet_columns(worksheet: Any) -> None:
    expected_columns = {
        0: "nominativo",
        4: "codice fiscale",
        5: "partita iva",
        6: "tipologia",
        7: "ultima mod",
        12: "utente",
    }
    header_row = next(
        worksheet.iter_rows(min_row=11, max_row=11, values_only=True),
        None,
    )
    if header_row is None:
        raise ValueError(REGISTRY_EXPECTED_FILE_DESCRIPTION)

    missing_columns = []
    for column_index, expected_label in expected_columns.items():
        actual_label = _normalize_header_label(
            header_row[column_index] if column_index < len(header_row) else None
        )
        if expected_label not in actual_label:
            missing_columns.append(expected_label)

    if missing_columns:
        raise ValueError(REGISTRY_EXPECTED_FILE_DESCRIPTION)


def _is_app_registry_worksheet(worksheet: Any) -> bool:
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        return False

    normalized_headers = [_normalize_header_label(value) for value in header_row]
    expected_headers = [_normalize_header_label(value) for value in APP_REGISTRY_EXCEL_HEADERS]
    return normalized_headers[: len(expected_headers)] == expected_headers


def _read_app_registry_records_from_worksheet(worksheet: Any) -> list[RegistryRecord]:
    records = []
    for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_index == 1:
            continue

        values = list(row[: len(APP_REGISTRY_EXCEL_HEADERS)])
        values.extend([None] * (len(APP_REGISTRY_EXCEL_HEADERS) - len(values)))
        if not any(_cell_text(value) for value in values):
            continue

        codice_anagrafica = _cell_text(values[0]) or None
        nominativo = _cell_text(values[1])
        normalized_name = normalize_registry_name(nominativo)
        codice_fiscale_raw = _cell_text(values[2])
        codice_fiscale = _normalized_registry_tax_code(codice_fiscale_raw)
        partita_iva_raw = _cell_text(values[3])
        partita_iva = _normalized_registry_vat_number(partita_iva_raw)

        records.append(
            RegistryRecord(
                registry_id=codice_anagrafica or f"row-{row_index}",
                source_row=row_index,
                codice_anagrafica=codice_anagrafica or extract_registry_code(nominativo),
                nominativo=nominativo,
                normalized_name=normalized_name,
                codice_fiscale_raw=codice_fiscale_raw,
                codice_fiscale=codice_fiscale,
                partita_iva_raw=partita_iva_raw,
                partita_iva=partita_iva,
                tipologia=_cell_text(values[4]),
                ultima_modifica=_json_text(values[5]) or None,
                utente=_cell_text(values[6]) or None,
                warnings=_registry_record_warnings(
                    normalized_name=normalized_name,
                    codice_fiscale_raw=codice_fiscale_raw,
                    codice_fiscale=codice_fiscale,
                    partita_iva_raw=partita_iva_raw,
                    partita_iva=partita_iva,
                ),
            )
        )

    return records


def _normalize_header_label(value: Any) -> str:
    text = _cell_text(value).lower()
    text = re.sub(r"[^0-9a-z]+", " ", text)
    return " ".join(text.split())


def _int_cell_value(value: Any, fallback: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return fallback


def _normalized_registry_tax_code(value: str) -> str | None:
    normalized = normalize_italian_tax_code(value)
    if normalized in {None, "0", "00000000000", "000000000000"}:
        return None
    return normalized


def _normalized_registry_vat_number(value: str) -> str | None:
    normalized = normalize_italian_vat_number(value)
    if normalized in {None, "00000000000"}:
        return None
    return normalized


def _registry_id_from_nominativo(nominativo: str, row_index: int) -> str:
    match = re.match(r"^(\d{3,})\b", nominativo.strip())
    if match:
        return match.group(1)

    return f"row-{row_index}"


def _registry_record_warnings(
    normalized_name: str,
    codice_fiscale_raw: str,
    codice_fiscale: str | None,
    partita_iva_raw: str,
    partita_iva: str | None,
) -> list[str]:
    warnings = []
    if not normalized_name:
        warnings.append("missing_normalized_name")
    if codice_fiscale_raw and codice_fiscale is None:
        warnings.append("invalid_codice_fiscale")
    if partita_iva_raw and partita_iva is None:
        warnings.append("invalid_partita_iva")

    return warnings


def _cell_text(value: Any) -> str:
    if value is None:
        return ""

    return str(value).strip()


def _json_text(value: Any) -> str | None:
    if value is None:
        return None

    isoformat = getattr(value, "isoformat", None)
    if callable(isoformat):
        return isoformat()

    text = _cell_text(value)
    return text or None


def _cache_to_jsonable(cache: RegistryCache) -> dict:
    return {
        "schema_version": cache.schema_version,
        "source_file": cache.source_file,
        "records": [asdict(record) for record in cache.records],
    }
