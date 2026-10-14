"""Durable, condominium-grouped accounting deliveries."""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import re
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from estrattore.invoice_registry_storage import (
    REGISTRY_LOCK, _field_values, _parse_money_value, _write_registry, load_invoice_registry,
)
from estrattore.registry import load_registry_cache
from estrattore.runtime_paths import get_runtime_paths


COLUMNS = [
    ("client_code", "Cod. cliente", 17),
    ("nome_emittente", "Denominazione", 42),
    ("data_fattura", "Data doc.", 15),
    ("numero_fattura", "N. doc", 18),
    ("totale_fattura", "Compenso", 18),
    ("codice_fiscale_emittente", "C.F. emittente", 23),
    ("p_iva", "P.IVA emittente", 19),
    ("ritenuta", "Ritenuta", 16),
    ("totale_netto_fattura", "Totale netto", 18),
    ("pagata_il", "Pagata il", 15),
]


def _normal(value):
    return " ".join(str(value or "").upper().split())


def _code(invoice, values, records, condominium=False):
    kinds = {"codice_fiscale", "condominio"} if condominium else {
        "codice_fiscale_emittente", "nome_emittente", "p_iva",
    }
    fields = [f for f in invoice.get("fields", []) if f.get("kind") in kinds]
    linked = {(f.get("registry_id"), f.get("registry_source_row")) for f in fields
              if f.get("registry_id")}
    matches = [records["linked"][key] for key in linked if key in records["linked"]]
    cf = _normal(values.get("codice_fiscale" if condominium else "codice_fiscale_emittente"))
    name = _normal(values.get("condominio" if condominium else "nome_emittente"))
    # Stale row references must not silently supply a different party's code.
    matches = [r for r in matches if (cf and _normal(r.codice_fiscale) == cf)
               or (name and _normal(r.nominativo) == name)]
    if not matches:
        matches = records["fiscal"].get(cf, []) if cf else []
        if not matches and name:
            matches = records["name"].get(name, [])
        if len(matches) > 1 and name:
            matches = [r for r in matches if _normal(r.nominativo) == name]
    codes = {r.codice_anagrafica for r in matches}
    return next(iter(codes)) if len(codes) == 1 else ""


def _safe_name(value, used, limit):
    name = re.sub(r'[\x00-\x1f\\/*?:\[\]<>|\"]', " ", value).strip(" .' ") or "Condominio"
    name = name[:limit]
    base = name
    suffix = 1
    while name.casefold() in used:
        suffix += 1
        tail = f" ({suffix})"
        name = base[:limit - len(tail)] + tail
    used.add(name.casefold())
    return name


def _text(cell, value):
    cell.value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))
    cell.data_type = "s"


def _date(value):
    for pattern in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value)[:10], pattern)
        except ValueError:
            pass
    return None


def _sheet(workbook, title, group, registry_records):
    sheet = workbook.create_sheet(title)
    first = group[0]
    values = _field_values(first.get("invoice", {}))
    code = _code(first.get("invoice", {}), values, registry_records, True)
    last_column = get_column_letter(len(COLUMNS))
    sheet.merge_cells(f"A1:{last_column}1")
    _text(sheet["A1"], values.get("condominio") or "Condominio non indicato")
    sheet["A1"].font = Font(name="Calibri", size=15, bold=True, color="205B45")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet[f"A1:{last_column}1"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="E7F1EA")
    sheet.row_dimensions[1].height = 38
    for column, value in enumerate(("Cod. anagrafica", code, "C.F. condominio", values.get("codice_fiscale")), 1):
        _text(sheet.cell(2, column), value)
    sheet.row_dimensions[2].height = 24
    for cell in sheet[2][:4]:
        cell.font = Font(name="Calibri", size=10, bold=cell.column in (1, 3), color="426554")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[3].height = 9
    for col, (_, label, width) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
        _text(sheet.cell(4, col), label)
    ordered = sorted(group, key=lambda r: (
        _date(_field_values(r.get("invoice", {})).get("data_fattura")) or datetime.min,
        _field_values(r.get("invoice", {})).get("numero_fattura", ""),
        r["record_id"],
    ))
    for row_number, record in enumerate(ordered, 5):
        invoice = record.get("invoice", {})
        row = _field_values(invoice)
        row["client_code"] = _code(invoice, row, registry_records)
        for col, (key, _, _) in enumerate(COLUMNS, 1):
            cell = sheet.cell(row_number, col)
            value = row.get(key, "")
            _text(cell, value)
            if key in {"totale_fattura", "ritenuta", "totale_netto_fattura"} and value:
                amount = _parse_money_value(value)
                if amount is not None:
                    cell.value = float(amount)
                    cell.number_format = '#,##0.00'
            elif key in {"data_fattura", "pagata_il"} and _date(value):
                cell.value = _date(value)
                cell.number_format = "dd/mm/yyyy"
        sheet.row_dimensions[row_number].height = 30
    for row in sheet.iter_rows(min_row=4, max_row=sheet.max_row, max_col=len(COLUMNS)):
        for cell in row:
            primary = cell.column <= 5
            header = cell.row == 4
            color = ("E0EDE4" if primary else "EDF1EE") if header else ("FFFFFF" if primary else "F7F8F7")
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(name="Calibri", size=10 if header else 11, bold=header,
                             color="205B45" if primary else "58655E")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="E8EDE9"))
    sheet.row_dimensions[4].height = 32
    sheet.sheet_view.showGridLines = False
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{last_column}{sheet.max_row}"
    sheet.oddFooter.center.text = "Pagina &P di &N"


def _workbook(groups, registry_records):
    workbook = Workbook()
    workbook.remove(workbook.active)
    used = set()
    for group in groups:
        name = _field_values(group[0].get("invoice", {})).get("condominio") or "Condominio non indicato"
        _sheet(workbook, _safe_name(name, used, 31), group, registry_records)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _withholding_summary(groups, registry_records):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Riepilogo ritenute"
    sheet.merge_cells("A1:E1")
    _text(sheet["A1"], "Riepilogo ritenute delle fatture esportate")
    sheet.row_dimensions[1].height = 36
    headers = ("Cod. anagrafica", "Condominio", "C.F. condominio", "Fatture", "Ritenute")
    for col, label in enumerate(headers, 1):
        _text(sheet.cell(2, col), label)
    total = Decimal("0")
    missing_total = 0
    invoice_count = 0
    for row_number, group in enumerate(groups, 3):
        invoice = group[0].get("invoice", {})
        values = _field_values(invoice)
        for col, value in enumerate((_code(invoice, values, registry_records, True),
                                      values.get("condominio") or "Condominio non indicato",
                                      values.get("codice_fiscale")), 1):
            _text(sheet.cell(row_number, col), value)
        amounts = [_parse_money_value(_field_values(r.get("invoice", {})).get("ritenuta", "")) for r in group]
        valid = [amount for amount in amounts if amount is not None and amount.is_finite() and amount >= 0]
        subtotal = sum(valid, Decimal("0"))
        missing = len(group) - len(valid)
        sheet.cell(row_number, 4, len(group))
        if valid:
            sheet.cell(row_number, 5, float(subtotal)).number_format = '#,##0.00'
        total += subtotal
        missing_total += missing
        invoice_count += len(group)
        if missing:
            for cell in sheet[row_number]:
                cell.fill = PatternFill("solid", fgColor="FFF5D6")
    last = len(groups) + 3
    _text(sheet.cell(last, 2), "Totale ritenute")
    sheet.cell(last, 4, invoice_count)
    if missing_total < invoice_count:
        sheet.cell(last, 5, float(total)).number_format = '#,##0.00'
    for row in sheet.iter_rows():
        for cell in row:
            emphasized = cell.row in (1, 2, last)
            cell.font = Font(name="Calibri", size=15 if cell.row == 1 else 11,
                             bold=emphasized, color="205B45" if emphasized else "36443D")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if emphasized:
                cell.fill = PatternFill("solid", fgColor="E7F1EA")
        if row[0].row >= 2:
            sheet.row_dimensions[row[0].row].height = 36
    for col, width in enumerate((19, 48, 24, 12, 23), 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:E{last}"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def export_invoices(record_ids, mode, runtime_paths=None, *, history_id=None, include_summary=False, deliver=None):
    if not isinstance(include_summary, bool):
        raise ValueError("Selezione del riepilogo non valida.")
    if mode not in {"workbook", "separate"}:
        raise ValueError("Scegli un formato di esportazione valido.")
    paths = runtime_paths or get_runtime_paths()
    with REGISTRY_LOCK:
        registry = load_invoice_registry(paths)
        if history_id is not None:
            previous_export = next((item for item in registry.get("export_history", [])
                                    if item.get("export_id") == history_id), None)
            if previous_export is None:
                raise ValueError("Esportazione non trovata nello storico.")
            record_ids = [item["record_id"] for item in previous_export["invoices"]]
        if not isinstance(record_ids, list) or not record_ids or not all(isinstance(i, str) for i in record_ids):
            raise ValueError("Seleziona almeno una fattura da esportare.")
        ids = set(record_ids)
        records = [r for r in registry["records"] if r.get("record_id") in ids
                   and r.get("status") == "saved" and r.get("registry_visible") is not False]
        if len(records) != len(ids):
            if history_id is not None:
                raise ValueError("Alcune fatture di questa esportazione non sono piu disponibili nel registro. Non e possibile riesportare il gruppo.")
            raise ValueError("Alcune fatture non sono piu disponibili. Aggiorna il registro.")
        cache_records = load_registry_cache(paths.registry_cache_file).records if paths.registry_cache_file.exists() else []
        registry_records = {"linked": {}, "fiscal": defaultdict(list), "name": defaultdict(list)}
        for record in cache_records:
            registry_records["linked"][(record.registry_id, record.source_row)] = record
            registry_records["fiscal"][_normal(record.codice_fiscale)].append(record)
            registry_records["name"][_normal(record.nominativo)].append(record)
        grouped = defaultdict(list)
        for record in records:
            fields = _field_values(record.get("invoice", {}))
            cf = _normal(fields.get("codice_fiscale"))
            name = _normal(fields.get("condominio"))
            key = ("cf", cf) if cf else (("name", name) if name else ("unknown", record["record_id"]))
            grouped[key].append(record)
        groups = [grouped[key] for key in sorted(grouped)]
        now = datetime.now().astimezone().isoformat(timespec="microseconds")
        export_id = uuid4().hex
        if mode == "workbook" and not include_summary:
            body = _workbook(groups, registry_records)
            name = "fatture_condomini.xlsx"
        else:
            output = BytesIO()
            used = set()
            with ZipFile(output, "w", ZIP_DEFLATED) as archive:
                if mode == "workbook":
                    archive.writestr("fatture_condomini.xlsx", _workbook(groups, registry_records))
                else:
                    for group in groups:
                        values = _field_values(group[0].get("invoice", {}))
                        label = f"Fatture {values.get('condominio') or 'Condominio non indicato'}"
                        filename = _safe_name(label, used, 100) + ".xlsx"
                        archive.writestr(filename, _workbook([group], registry_records))
                if include_summary:
                    archive.writestr("riepilogo_ritenute.xlsx", _withholding_summary(groups, registry_records))
            body = output.getvalue()
            name = "fatture_condomini.zip"
        if deliver is not None:
            deliver(body, name)
        directory = paths.archive_dir / "exports" / export_id
        directory.mkdir(parents=True)
        target = directory / name
        try:
            target.write_bytes(body)
            url = f"/runtime/archive/exports/{export_id}/{name}"
            for record in records:
                record["last_exported_at"] = now
                record["last_export_url"] = url
            registry.setdefault("export_history", []).append({
                "export_id": export_id,
                "exported_at": now,
                "mode": mode,
                "include_summary": include_summary,
                "download_url": url,
                "source_export_id": history_id,
                "invoices": [
                    {
                        "record_id": record["record_id"],
                        **{key: _field_values(record.get("invoice", {})).get(key, "")
                           for key in ("numero_fattura", "data_fattura", "nome_emittente",
                                       "condominio", "totale_fattura")},
                    }
                    for record in records
                ],
            })
            registry["updated_at"] = now
            _write_registry(paths.invoice_registry_file, registry)
        except Exception:
            target.unlink(missing_ok=True)
            directory.rmdir()
            raise
        return {"ok": True, "export_id": export_id, "download_url": url, "file_name": name,
                "invoice_count": len(records), "condominium_count": len(groups),
                "exported_at": now}


def delete_export_history(export_id, runtime_paths=None):
    """Remove a receipt, preserving invoices and still-referenced downloads."""
    paths = runtime_paths or get_runtime_paths()
    with REGISTRY_LOCK:
        registry = load_invoice_registry(paths)
        history = registry.get("export_history", [])
        target = next((item for item in history if item.get("export_id") == export_id), None)
        if target is None:
            raise ValueError("Esportazione non trovata nello storico.")
        registry["export_history"] = [item for item in history if item.get("export_id") != export_id]
        registry["updated_at"] = datetime.now().astimezone().isoformat()
        _write_registry(paths.invoice_registry_file, registry)
        url = target.get("download_url", "")
        referenced = any(r.get("last_export_url") == url for r in registry["records"])
        referenced = referenced or any(item.get("download_url") == url for item in registry["export_history"])
        if not referenced and isinstance(export_id, str) and re.fullmatch(r"[0-9a-f]{32}", export_id):
            directory = paths.archive_dir / "exports" / export_id
            if directory.resolve().parent == (paths.archive_dir / "exports").resolve():
                try:
                    for filename in ("fatture_condomini.xlsx", "fatture_condomini.zip"):
                        (directory / filename).unlink(missing_ok=True)
                    if directory.exists():
                        directory.rmdir()
                except OSError:
                    import logging
                    logging.getLogger(__name__).warning("Unable to remove unused export artifact %s", export_id, exc_info=True)
        return {"ok": True}
