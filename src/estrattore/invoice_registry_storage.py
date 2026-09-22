from __future__ import annotations

import json
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from estrattore.feedback_collection import (
    commit_review_operation,
    feedback_operation_id,
    mark_archive_paths_unavailable,
    prepare_registry_edit_operation,
    prepare_review_operation,
    reconcile_feedback_operations,
    record_capture_issue,
)
from estrattore.invoice_pdf_archive import archive_invoice_pdf, invoice_page_label
from estrattore.processing_storage import load_review_invoice_ref, mark_review_invoice
from estrattore.runtime_paths import RuntimePaths, get_runtime_paths


INVOICE_REGISTRY_SCHEMA_VERSION = 3
INVOICE_EXPORT_COLUMNS = [
    ("saved_at", "Salvata il"),
    ("nome_emittente", "Nome emittente"),
    ("codice_fiscale_emittente", "C.F. emittente"),
    ("p_iva", "P.IVA emittente"),
    ("condominio", "Condominio"),
    ("codice_fiscale", "C.F. condominio"),
    ("numero_fattura", "Numero fattura"),
    ("data_fattura", "Data fattura"),
    ("totale_fattura", "Totale fattura"),
    ("ritenuta", "Ritenuta"),
    ("totale_netto_fattura", "Totale netto fattura"),
    ("pagata_il", "Data pagamento"),
    ("source_file_name", "File origine"),
    ("archived_pdf_path", "PDF archiviato"),
    ("archived_source_pdf_path", "PDF sorgente"),
]

INVOICE_FIELD_ALIASES = {
    "totale_fattura": ["totale_fattura", "economic_summary.totale_fattura.value"],
    "ritenuta": ["ritenuta", "economic_summary.ritenuta.value"],
    "totale_netto_fattura": [
        "totale_netto_fattura",
        "economic_summary.totale_netto_fattura.value",
    ],
}
WITHHOLDING_FIELD_KINDS = {"ritenuta", "economic_summary.ritenuta.value"}
ACCOUNTING_FIELD_KINDS = {
    kind
    for aliases in INVOICE_FIELD_ALIASES.values()
    for kind in aliases
}


def load_invoice_registry(
    runtime_paths: RuntimePaths | None = None,
) -> dict[str, Any]:
    paths = runtime_paths or get_runtime_paths()
    if not paths.invoice_registry_file.exists():
        return _empty_registry()

    payload = json.loads(paths.invoice_registry_file.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Registro fatture non valido")

    try:
        stored_schema_version = int(payload.get("schema_version") or 1)
    except (TypeError, ValueError):
        stored_schema_version = 1
    payload.setdefault("updated_at", "")
    payload.setdefault("records", [])
    payload["records"] = [
        record for record in payload["records"] if isinstance(record, dict)
    ]
    if stored_schema_version < 3:
        payload["records"] = [
            _migrate_legacy_withholding_record(record)
            for record in payload["records"]
        ]
    payload["schema_version"] = INVOICE_REGISTRY_SCHEMA_VERSION
    return payload


def purge_discarded_invoice_records(
    runtime_paths: RuntimePaths | None = None,
) -> int:
    paths = runtime_paths or get_runtime_paths()
    if not paths.invoice_registry_file.is_file():
        return 0
    raw_payload = json.loads(paths.invoice_registry_file.read_text(encoding="utf-8"))
    records = raw_payload.get("records", []) if isinstance(raw_payload, dict) else []
    try:
        schema_version = int(raw_payload.get("schema_version") or 1)
    except (TypeError, ValueError):
        schema_version = 1
    next_records = [
        record
        for record in records
        if isinstance(record, dict)
        and not (
            schema_version < INVOICE_REGISTRY_SCHEMA_VERSION
            and record.get("status") == "discarded"
            and not str(record.get("archived_pdf_path") or "")
        )
    ]
    if schema_version < 3:
        next_records = [
            _migrate_legacy_withholding_record(record)
            for record in next_records
        ]
    removed = len(records) - len(next_records)
    if removed or schema_version < INVOICE_REGISTRY_SCHEMA_VERSION:
        raw_payload["records"] = next_records
        raw_payload["schema_version"] = INVOICE_REGISTRY_SCHEMA_VERSION
        raw_payload["updated_at"] = _timestamp()
        _write_registry(paths.invoice_registry_file, raw_payload)
    return removed


def save_reviewed_invoice(
    payload: dict[str, Any],
    runtime_paths: RuntimePaths | None = None,
    now: datetime | None = None,
) -> dict[str, Any]:
    paths = runtime_paths or get_runtime_paths()
    paths.ensure_directories()

    invoice_payload = _invoice_payload(payload)
    _ensure_valid_accounting_fields(invoice_payload)
    review_batch_base_path = _required_string(payload, "review_batch_base_path")
    record_id = _record_id(invoice_payload, review_batch_base_path)
    raw_invoice_id = _raw_invoice_id(invoice_payload)
    registry = load_invoice_registry(paths)
    existing_record = next(
        (
            record
            for record in registry.get("records", [])
            if isinstance(record, dict)
            and str(record.get("record_id") or "") == record_id
            and record.get("status") == "saved"
        ),
        None,
    )
    if existing_record is not None:
        try:
            reconcile_feedback_operations([existing_record], paths)
        except Exception as error:
            _record_feedback_issue(
                process_id=str(existing_record.get("source_batch_id") or ""),
                document_id=str(existing_record.get("document_id") or ""),
                stage="invoice_save_recovery",
                error=error,
                runtime_paths=paths,
            )
        try:
            invoice_ref = load_review_invoice_ref(
                review_batch_base_path, raw_invoice_id, paths
            )
        except (FileNotFoundError, ValueError):
            return existing_record
        operation_id = _prepare_feedback_review_operation(
            paths=paths,
            invoice_payload=(
                existing_record.get("invoice")
                if isinstance(existing_record.get("invoice"), dict)
                else invoice_payload
            ),
            invoice_ref=invoice_ref,
            review_batch_base_path=review_batch_base_path,
            record_id=record_id,
            event_kind="invoice_save",
            outcome="saved",
        )
        _commit_feedback_review_operation(
            operation_id, paths, invoice_ref, "invoice_save"
        )
        mark_review_invoice(
            review_batch_base_path, raw_invoice_id, "saved", paths
        )
        return existing_record

    invoice_ref = load_review_invoice_ref(review_batch_base_path, raw_invoice_id, paths)
    operation_id = _prepare_feedback_review_operation(
        paths=paths,
        invoice_payload=invoice_payload,
        invoice_ref=invoice_ref,
        review_batch_base_path=review_batch_base_path,
        record_id=record_id,
        event_kind="invoice_save",
        outcome="saved",
    )
    parent_invoice_record = _parent_invoice_record(invoice_ref, registry)
    archived_pdf = archive_invoice_pdf(
        invoice_payload,
        invoice_ref,
        paths,
        parent_invoice_record=parent_invoice_record,
    )
    saved_at = _timestamp(now)
    registry_record = {
        "record_id": record_id,
        "status": "saved",
        "saved_at": saved_at,
        "updated_at": saved_at,
        "review_batch_base_path": review_batch_base_path,
        "selected_batch_id": str(payload.get("selected_batch_id") or ""),
        "source_batch_id": invoice_ref.source_batch_id,
        "document_id": invoice_ref.document_id,
        "source_file_name": invoice_ref.source_file_name,
        "archived_pdf_path": archived_pdf.relative_path,
        "archived_pdf_url": f"/runtime/{archived_pdf.relative_path}",
        "archived_pdf_page_numbers": archived_pdf.page_numbers,
        "archived_source_pdf_path": archived_pdf.archived_source_relative_path,
        "archived_source_pdf_url": (
            f"/runtime/{archived_pdf.archived_source_relative_path}"
            if archived_pdf.archived_source_relative_path
            else ""
        ),
        "invoice": invoice_payload,
    }

    registry["records"] = _upsert_record(
        registry.get("records", []),
        registry_record,
    )
    registry["updated_at"] = saved_at
    _write_registry(paths.invoice_registry_file, registry)
    _commit_feedback_review_operation(
        operation_id, paths, invoice_ref, "invoice_save"
    )
    mark_review_invoice(
        review_batch_base_path,
        raw_invoice_id,
        "saved",
        paths,
    )
    return registry_record


def discard_reviewed_invoice(
    payload: dict[str, Any],
    runtime_paths: RuntimePaths | None = None,
    now: datetime | None = None,
) -> dict[str, Any]:
    paths = runtime_paths or get_runtime_paths()
    paths.ensure_directories()

    invoice_payload = _invoice_payload(payload)
    review_batch_base_path = _required_string(payload, "review_batch_base_path")
    raw_invoice_id = _raw_invoice_id(invoice_payload)
    record_id = _record_id(invoice_payload, review_batch_base_path)
    registry = load_invoice_registry(paths)
    existing_record = next(
        (
            record
            for record in registry.get("records", [])
            if isinstance(record, dict)
            and str(record.get("record_id") or "") == record_id
            and record.get("status") == "discarded"
            and str(record.get("archived_pdf_path") or "")
        ),
        None,
    )
    if existing_record is not None:
        try:
            reconcile_feedback_operations([existing_record], paths)
        except Exception as error:
            _record_feedback_issue(
                process_id=str(existing_record.get("source_batch_id") or ""),
                document_id=str(existing_record.get("document_id") or ""),
                stage="invoice_discard_recovery",
                error=error,
                runtime_paths=paths,
            )
        try:
            invoice_ref = load_review_invoice_ref(
                review_batch_base_path, raw_invoice_id, paths
            )
        except (FileNotFoundError, ValueError):
            return existing_record
        operation_id = _prepare_feedback_review_operation(
            paths=paths,
            invoice_payload=(
                existing_record.get("invoice")
                if isinstance(existing_record.get("invoice"), dict)
                else invoice_payload
            ),
            invoice_ref=invoice_ref,
            review_batch_base_path=review_batch_base_path,
            record_id=record_id,
            event_kind="invoice_discard",
            outcome="discarded",
        )
        _commit_feedback_review_operation(
            operation_id, paths, invoice_ref, "invoice_discard"
        )
        mark_review_invoice(
            review_batch_base_path, raw_invoice_id, "discarded", paths
        )
        return existing_record

    invoice_ref = load_review_invoice_ref(review_batch_base_path, raw_invoice_id, paths)
    operation_id = _prepare_feedback_review_operation(
        paths=paths,
        invoice_payload=invoice_payload,
        invoice_ref=invoice_ref,
        review_batch_base_path=review_batch_base_path,
        record_id=record_id,
        event_kind="invoice_discard",
        outcome="discarded",
    )
    parent_invoice_record = _parent_invoice_record(invoice_ref, registry)
    local_page_numbers = _page_numbers_payload(invoice_payload)
    source_page_numbers = _compose_source_page_numbers(
        local_page_numbers,
        invoice_ref.source_page_numbers,
    )
    display_name = invoice_page_label(
        invoice_ref.source_file_name,
        source_page_numbers,
    )
    archived_pdf = archive_invoice_pdf(
        invoice_payload,
        invoice_ref,
        paths,
        parent_invoice_record=parent_invoice_record,
        file_stem=display_name,
    )
    discarded_at = _timestamp(now)
    result = {
        "record_id": record_id,
        "status": "discarded",
        "discarded_at": discarded_at,
        "updated_at": discarded_at,
        "review_batch_base_path": review_batch_base_path,
        "selected_batch_id": str(payload.get("selected_batch_id") or ""),
        "source_batch_id": invoice_ref.source_batch_id,
        "document_id": invoice_ref.document_id,
        "source_file_name": invoice_ref.source_file_name,
        "display_name": display_name,
        "archived_pdf_path": archived_pdf.relative_path,
        "archived_pdf_url": f"/runtime/{archived_pdf.relative_path}",
        "archived_pdf_page_numbers": archived_pdf.page_numbers,
        "archived_source_pdf_path": archived_pdf.archived_source_relative_path,
        "archived_source_pdf_url": (
            f"/runtime/{archived_pdf.archived_source_relative_path}"
            if archived_pdf.archived_source_relative_path
            else ""
        ),
        "invoice": invoice_payload,
    }
    registry["records"] = _upsert_record(registry.get("records", []), result)
    registry["updated_at"] = discarded_at
    _write_registry(paths.invoice_registry_file, registry)
    _commit_feedback_review_operation(
        operation_id, paths, invoice_ref, "invoice_discard"
    )
    mark_review_invoice(
        review_batch_base_path,
        raw_invoice_id,
        "discarded",
        paths,
    )
    return result


def save_invoice_registry_excel(
    output_path: Path | None = None,
    runtime_paths: RuntimePaths | None = None,
) -> Path:
    paths = runtime_paths or get_runtime_paths()
    registry = load_invoice_registry(paths)
    target_path = output_path or paths.invoice_registry_export_xlsx_file
    records = [
        record
        for record in registry.get("records", [])
        if (
            isinstance(record, dict)
            and record.get("status") == "saved"
            and record.get("registry_visible") is not False
        )
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registro fatture"
    header_fill = PatternFill("solid", fgColor="DDEFE6")
    header_font = Font(bold=True, color="174C3A")
    for column_index, (_, label) in enumerate(INVOICE_EXPORT_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=label)
        cell.fill = header_fill
        cell.font = header_font

    for row_index, record in enumerate(records, start=2):
        row_values = _invoice_export_row(record)
        for column_index, (column_key, _) in enumerate(INVOICE_EXPORT_COLUMNS, start=1):
            sheet.cell(row=row_index, column=column_index, value=row_values.get(column_key, ""))

    for column_cells in sheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length + 2, 14),
            44,
        )
    sheet.freeze_panes = "A2"
    target_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target_path)
    return target_path


def save_invoice_registry_records(
    records: list[dict[str, Any]],
    runtime_paths: RuntimePaths | None = None,
    now: datetime | None = None,
) -> dict[str, Any]:
    paths = runtime_paths or get_runtime_paths()
    registry = load_invoice_registry(paths)
    existing_records = [
        record
        for record in registry.get("records", [])
        if isinstance(record, dict)
    ]
    try:
        reconcile_feedback_operations(existing_records, paths)
    except Exception as error:
        _record_feedback_issue(
            process_id="",
            document_id="",
            stage="registry_edit_recovery",
            error=error,
            runtime_paths=paths,
        )
    existing_by_id = {
        str(record.get("record_id") or ""): record
        for record in existing_records
        if str(record.get("record_id") or "")
    }
    updated_at = _timestamp(now)
    next_records: list[dict[str, Any]] = []
    feedback_operation_ids: list[str] = []

    for record_payload in records:
        if not isinstance(record_payload, dict):
            continue
        record_id = str(record_payload.get("record_id") or "").strip()
        if not record_id or record_id not in existing_by_id:
            continue

        existing_record = existing_by_id[record_id]
        if existing_record.get("status") != "saved":
            next_records.append(existing_record)
            continue

        updated_invoice = _updated_invoice_payload(
            existing_record.get("invoice"),
            record_payload.get("invoice"),
        )
        _ensure_valid_accounting_fields(updated_invoice)
        next_record = {
            **existing_record,
            "updated_at": updated_at,
            "invoice": updated_invoice,
        }
        if _invoice_field_values(existing_record.get("invoice")) != _invoice_field_values(
            updated_invoice
        ):
            operation_id = _prepare_registry_feedback_operation(
                paths=paths,
                record_id=record_id,
                previous_invoice_payload=(
                    existing_record.get("invoice")
                    if isinstance(existing_record.get("invoice"), dict)
                    else {}
                ),
                updated_invoice_payload=updated_invoice,
                legacy_record=existing_record,
            )
            if operation_id:
                feedback_operation_ids.append(operation_id)
        next_records.append(next_record)

    submitted_record_ids = {
        str(record.get("record_id") or "")
        for record in records
        if isinstance(record, dict)
    }
    next_records.extend(
        record
        for record in existing_records
        if str(record.get("record_id") or "") not in submitted_record_ids
    )

    registry["records"] = next_records
    registry["updated_at"] = updated_at
    _write_registry(paths.invoice_registry_file, registry)
    for operation_id in feedback_operation_ids:
        try:
            commit_review_operation(operation_id, paths)
        except Exception as error:
            _record_feedback_issue(
                process_id="",
                document_id="",
                stage="registry_edit_commit",
                error=error,
                runtime_paths=paths,
            )
    try:
        reconcile_feedback_operations(
            registry["records"],
            paths,
            abandon_mismatched=True,
        )
    except Exception as error:
        _record_feedback_issue(
            process_id="",
            document_id="",
            stage="registry_edit_recovery",
            error=error,
            runtime_paths=paths,
        )
    return registry


def delete_invoice_archive(
    scope: str,
    record_id: str | None = None,
    record_ids: list[str] | None = None,
    archived_source_pdf_path: str | None = None,
    archived_source_pdf_paths: list[str] | None = None,
    runtime_paths: RuntimePaths | None = None,
) -> dict[str, Any]:
    """Delete one persisted projection without cascading into sibling archives."""
    paths = runtime_paths or get_runtime_paths()
    registry = load_invoice_registry(paths)
    persisted_records = [
        record
        for record in registry.get("records", [])
        if isinstance(record, dict) and record.get("status") in {"saved", "discarded"}
    ]

    if scope in {"invoice", "registry"}:
        target_ids = {
            str(value).strip()
            for value in [*(record_ids or []), record_id]
            if str(value or "").strip()
        }
        targets = [
            record
            for record in persisted_records
            if str(record.get("record_id") or "") in target_ids
        ]
        if not targets:
            raise ValueError("Fatture non trovate")
    elif scope == "source":
        target_sources = {
            str(value).strip()
            for value in [*(archived_source_pdf_paths or []), archived_source_pdf_path]
            if str(value or "").strip()
        }
        if not target_sources:
            raise ValueError("File origine archiviato mancante")
        targets = [
            record
            for record in persisted_records
            if str(record.get("archived_source_pdf_path") or "") in target_sources
        ]
        if not targets:
            raise ValueError("File origine archiviati non trovati")
    elif scope == "all":
        targets = persisted_records
    else:
        raise ValueError("Ambito archivio non valido")

    target_ids = {str(record.get("record_id") or "") for record in targets}
    target_source_paths = {
        str(record.get("archived_source_pdf_path") or "")
        for record in targets
        if str(record.get("archived_source_pdf_path") or "")
    }
    all_archive_paths: set[str] = set()
    if scope == "all":
        for archive_root in (paths.archive_invoices_dir, paths.archive_sources_dir):
            all_archive_paths.update(
                candidate.relative_to(paths.app_data_dir).as_posix()
                for candidate in archive_root.rglob("*")
                if candidate.is_file()
            )
    deleted_invoice_file_count = 0
    deleted_source_file_count = 0
    deleted_registry_count = 0
    deleted_at = _timestamp()

    if scope in {"source", "all"}:
        for source_path in target_source_paths:
            if _remove_archive_file(
                source_path,
                paths.archive_sources_dir,
                paths.app_data_dir,
            ):
                deleted_source_file_count += 1

    next_records: list[dict[str, Any]] = []
    for record in registry.get("records", []):
        if not isinstance(record, dict):
            continue
        if str(record.get("record_id") or "") not in target_ids:
            next_records.append(record)
            continue

        next_record = dict(record)
        if scope == "registry":
            next_record["registry_visible"] = False
            next_record["registry_deleted_at"] = deleted_at
            deleted_registry_count += 1

        if scope in {"invoice", "all"}:
            if _remove_archive_file(
                record.get("archived_pdf_path"),
                paths.archive_invoices_dir,
                paths.app_data_dir,
            ):
                deleted_invoice_file_count += 1
            next_record["invoice_archive_visible"] = False
            next_record["archived_pdf_path"] = ""
            next_record["archived_pdf_url"] = ""

        if scope in {"source", "all"}:
            next_record["source_archive_visible"] = False
            next_record["archived_source_pdf_path"] = ""
            next_record["archived_source_pdf_url"] = ""

        next_record["updated_at"] = deleted_at
        next_records.append(next_record)

    if scope == "all":
        deleted_invoice_file_count += _remove_all_archive_files(
            paths.archive_invoices_dir,
            paths.app_data_dir,
        )
        deleted_source_file_count += _remove_all_archive_files(
            paths.archive_sources_dir,
            paths.app_data_dir,
        )

    registry["records"] = next_records
    registry["updated_at"] = deleted_at
    _write_registry(paths.invoice_registry_file, registry)
    removed_archive_paths = set()
    if scope in {"invoice", "all"}:
        removed_archive_paths.update(
            str(record.get("archived_pdf_path") or "") for record in targets
        )
    if scope in {"source", "all"}:
        removed_archive_paths.update(target_source_paths)
    removed_archive_paths.update(all_archive_paths)
    try:
        mark_archive_paths_unavailable(removed_archive_paths, paths)
    except Exception as error:
        _record_feedback_issue(
            process_id="",
            document_id="",
            stage="archive_delete_availability",
            error=error,
            runtime_paths=paths,
        )
    return {
        "registry": registry,
        "deleted_invoice_count": len(targets),
        "deleted_registry_count": deleted_registry_count,
        "deleted_invoice_file_count": deleted_invoice_file_count,
        "deleted_source_file_count": deleted_source_file_count,
    }


def _remove_archive_file(
    relative_path: Any,
    allowed_directory: Path,
    app_data_dir: Path,
) -> bool:
    value = str(relative_path or "").strip()
    if not value:
        return False

    candidate = (app_data_dir / value).resolve()
    allowed_root = allowed_directory.resolve()
    if candidate == allowed_root or allowed_root not in candidate.parents:
        raise ValueError("Percorso archivio non autorizzato")
    if not candidate.is_file():
        return False
    candidate.unlink()
    return True


def _remove_all_archive_files(allowed_directory: Path, app_data_dir: Path) -> int:
    allowed_root = allowed_directory.resolve()
    if allowed_root == app_data_dir.resolve() or app_data_dir.resolve() not in allowed_root.parents:
        raise ValueError("Percorso archivio non autorizzato")

    deleted_count = 0
    for candidate in allowed_root.rglob("*"):
        if not candidate.is_file():
            continue
        resolved_candidate = candidate.resolve()
        if allowed_root not in resolved_candidate.parents:
            raise ValueError("Percorso archivio non autorizzato")
        resolved_candidate.unlink()
        deleted_count += 1
    return deleted_count


def _invoice_payload(payload: dict[str, Any]) -> dict[str, Any]:
    invoice = payload.get("invoice")
    if not isinstance(invoice, dict):
        raise ValueError("Fattura da salvare mancante")

    invoice_id = str(invoice.get("invoiceId") or "").strip()
    if not invoice_id:
        raise ValueError("Identificativo fattura mancante")

    fields = invoice.get("fields", [])
    if not isinstance(fields, list):
        raise ValueError("Campi fattura non validi")

    return {
        "invoiceId": invoice_id,
        "manifestInvoiceId": str(invoice.get("manifestInvoiceId") or ""),
        "title": str(invoice.get("title") or ""),
        "documentId": str(invoice.get("documentId") or ""),
        "sourceFile": str(invoice.get("sourceFile") or ""),
        "basePath": str(invoice.get("basePath") or ""),
        "pageNumbers": _page_numbers_payload(invoice),
        "pages": invoice.get("pages", []),
        "fields": [
            field
            for field in fields
            if isinstance(field, dict) and str(field.get("kind") or "").strip()
        ],
    }


def _updated_invoice_payload(
    existing_invoice: object,
    submitted_invoice: object,
) -> dict[str, Any]:
    if not isinstance(existing_invoice, dict):
        existing_invoice = {}
    if not isinstance(submitted_invoice, dict):
        return dict(existing_invoice)

    existing_fields = existing_invoice.get("fields", [])
    submitted_fields = submitted_invoice.get("fields", [])
    if not isinstance(existing_fields, list) or not isinstance(submitted_fields, list):
        return dict(existing_invoice)

    submitted_by_kind = {
        str(field.get("kind") or ""): field
        for field in submitted_fields
        if isinstance(field, dict) and str(field.get("kind") or "").strip()
    }
    next_fields = []
    existing_kinds: set[str] = set()
    for field in existing_fields:
        if not isinstance(field, dict):
            continue
        kind = str(field.get("kind") or "").strip()
        if not kind:
            continue
        existing_kinds.add(kind)
        submitted_field = submitted_by_kind.get(kind, {})
        next_fields.append(
            {
                **field,
                **{
                    key: submitted_field[key]
                    for key in (
                        "review_interaction",
                        "review_value_origin",
                        "review_action_id",
                        "review_source_field_kind",
                        "review_registry_snapshot",
                    )
                    if key in submitted_field
                },
                "value": str(submitted_field.get("value", field.get("value") or "")),
            }
        )

    for field in submitted_fields:
        if not isinstance(field, dict):
            continue
        kind = str(field.get("kind") or "").strip()
        if not kind or kind in existing_kinds:
            continue
        next_fields.append({
            **field,
            "kind": kind,
            "value": str(field.get("value") or ""),
        })

    return {
        **existing_invoice,
        "fields": next_fields,
    }


def _prepare_feedback_review_operation(
    *,
    paths: RuntimePaths,
    invoice_payload: dict[str, Any],
    invoice_ref: Any,
    review_batch_base_path: str,
    record_id: str,
    event_kind: str,
    outcome: str,
) -> str | None:
    operation_id = feedback_operation_id(
        event_kind=event_kind,
        record_id=record_id,
        invoice_payload=invoice_payload,
    )
    try:
        return prepare_review_operation(
            operation_id=operation_id,
            event_kind=event_kind,
            outcome=outcome,
            record_id=record_id,
            invoice_payload=invoice_payload,
            review_batch_base_path=review_batch_base_path,
            process_id=str(invoice_ref.process_id or ""),
            source_batch_id=str(invoice_ref.source_batch_id or ""),
            document_id=str(invoice_ref.document_id or ""),
            source_file_name=str(invoice_ref.source_file_name or ""),
            parent_invoice_record_id=str(invoice_ref.parent_invoice_record_id or ""),
            source_page_numbers=list(invoice_ref.source_page_numbers or []),
            runtime_paths=paths,
        )
    except Exception as error:
        _record_feedback_issue(
            process_id=str(invoice_ref.process_id or ""),
            document_id=str(invoice_ref.document_id or ""),
            stage=f"{event_kind}_prepare",
            error=error,
            runtime_paths=paths,
        )
        return None


def _commit_feedback_review_operation(
    operation_id: str | None,
    paths: RuntimePaths,
    invoice_ref: Any,
    event_kind: str,
) -> None:
    if not operation_id:
        return
    try:
        commit_review_operation(operation_id, paths)
    except Exception as error:
        _record_feedback_issue(
            process_id=str(invoice_ref.process_id or ""),
            document_id=str(invoice_ref.document_id or ""),
            stage=f"{event_kind}_commit",
            error=error,
            runtime_paths=paths,
        )


def _prepare_registry_feedback_operation(
    *,
    paths: RuntimePaths,
    record_id: str,
    previous_invoice_payload: dict[str, Any],
    updated_invoice_payload: dict[str, Any],
    legacy_record: dict[str, Any],
) -> str | None:
    try:
        return prepare_registry_edit_operation(
            operation_id=None,
            record_id=record_id,
            previous_invoice_payload=previous_invoice_payload,
            updated_invoice_payload=updated_invoice_payload,
            legacy_record=legacy_record,
            runtime_paths=paths,
        )
    except Exception as error:
        _record_feedback_issue(
            process_id="",
            document_id="",
            stage="registry_edit_prepare",
            error=error,
            runtime_paths=paths,
        )
        return None


def _record_feedback_issue(**kwargs: Any) -> None:
    try:
        record_capture_issue(**kwargs)
    except Exception:
        # Collection diagnostics cannot invalidate an accounting action.
        return


def _invoice_field_values(invoice: object) -> dict[str, str]:
    if not isinstance(invoice, dict):
        return {}
    fields = invoice.get("fields", [])
    if not isinstance(fields, list):
        return {}
    return {
        str(field.get("kind") or ""): str(field.get("value") or "")
        for field in fields
        if isinstance(field, dict) and str(field.get("kind") or "")
    }


def _parse_money_value(value: str) -> Decimal | None:
    compact = re.sub(r"[\s\u00a0€]", "", value.strip())
    if not compact or re.fullmatch(r"[+-]?[\d.,]+", compact) is None:
        return None

    sign = "-" if compact.startswith("-") else ""
    unsigned = compact.lstrip("+-")
    comma_index = unsigned.rfind(",")
    dot_index = unsigned.rfind(".")

    if comma_index >= 0 and dot_index >= 0:
        decimal_index = max(comma_index, dot_index)
        fraction = unsigned[decimal_index + 1 :]
        if not fraction.isdigit() or len(fraction) not in (1, 2):
            return None
        whole = re.sub(r"[.,]", "", unsigned[:decimal_index])
        normalized = f"{sign}{whole}.{fraction.ljust(2, '0')}"
    elif comma_index >= 0 or dot_index >= 0:
        separator = "," if comma_index >= 0 else "."
        parts = unsigned.split(separator)
        fraction = parts[-1]
        if len(parts) == 2 and fraction.isdigit() and len(fraction) in (1, 2):
            normalized = f"{sign}{parts[0]}.{fraction.ljust(2, '0')}"
        else:
            normalized = f"{sign}{''.join(parts)}.00"
    else:
        normalized = f"{sign}{unsigned}.00"

    try:
        return Decimal(normalized).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _invoice_amount_by_alias(
    field_values: dict[str, str],
    canonical_kind: str,
) -> Decimal | None:
    for kind in INVOICE_FIELD_ALIASES[canonical_kind]:
        value = field_values.get(kind, "").strip()
        if value:
            return _parse_money_value(value)
    return None


def _migrate_legacy_withholding_record(record: dict[str, Any]) -> dict[str, Any]:
    invoice = record.get("invoice")
    if not isinstance(invoice, dict) or not isinstance(invoice.get("fields"), list):
        return record

    field_values = _invoice_field_values(invoice)
    total = _invoice_amount_by_alias(field_values, "totale_fattura")
    net = _invoice_amount_by_alias(field_values, "totale_netto_fattura")
    derived = total - net if total is not None and net is not None and total >= net else None

    migrated_fields = []
    for field in invoice["fields"]:
        if not isinstance(field, dict):
            migrated_fields.append(field)
            continue
        kind = str(field.get("kind") or "").strip()
        if kind not in WITHHOLDING_FIELD_KINDS:
            migrated_fields.append(field)
            continue
        is_legacy_percentage = "%" in str(field.get("value") or "") or "%" in str(
            field.get("label") or ""
        )
        migrated_fields.append(
            {
                **field,
                "label": "Ritenuta",
                "value": f"{derived:.2f}" if is_legacy_percentage and derived is not None else (
                    "" if is_legacy_percentage else str(field.get("value") or "")
                ),
            }
        )

    return {
        **record,
        "invoice": {
            **invoice,
            "fields": migrated_fields,
        },
    }


def _ensure_valid_accounting_fields(invoice: dict[str, Any]) -> None:
    fields = invoice.get("fields", [])
    if not isinstance(fields, list):
        return

    for field in fields:
        if not isinstance(field, dict):
            continue
        if str(field.get("kind") or "").strip() not in ACCOUNTING_FIELD_KINDS:
            continue
        value = str(field.get("value") or "").strip()
        amount = _parse_money_value(value) if value else None
        if value and (
            value.startswith(("+", "-"))
            or amount is None
            or amount < 0
        ):
            label = str(field.get("label") or "Importo").strip() or "Importo"
            raise ValueError(
                f"Importo non valido ({label}): inserisci un valore numerico "
                "senza segni o lettere"
            )

    field_values = _invoice_field_values(invoice)
    total = _invoice_amount_by_alias(field_values, "totale_fattura")
    withholding = _invoice_amount_by_alias(field_values, "ritenuta")
    net = _invoice_amount_by_alias(field_values, "totale_netto_fattura")
    if (
        total is not None
        and withholding is not None
        and withholding >= 0
        and net is not None
        and abs((total - withholding) - net) > Decimal("0.01")
    ):
        raise ValueError(
            "Valori contabili non coerenti: totale fattura meno ritenuta "
            "deve corrispondere al totale netto fattura"
        )


def _record_id(invoice_payload: dict[str, Any], review_batch_base_path: str) -> str:
    return ":".join(
        [
            "invoice",
            review_batch_base_path.strip("/").replace("/", ":"),
            str(invoice_payload["invoiceId"]),
        ]
    )


def _raw_invoice_id(invoice_payload: dict[str, Any]) -> str:
    manifest_invoice_id = str(invoice_payload.get("manifestInvoiceId") or "").strip()
    if manifest_invoice_id:
        return manifest_invoice_id

    invoice_id = str(invoice_payload["invoiceId"])
    return invoice_id.rsplit(":", 1)[-1]


def _parent_invoice_record(
    invoice_ref: Any,
    registry: dict[str, Any],
) -> dict[str, Any] | None:
    parent_record_id = str(invoice_ref.parent_invoice_record_id or "")
    if not parent_record_id:
        return None
    for record in registry.get("records", []):
        if (
            isinstance(record, dict)
            and str(record.get("record_id") or "") == parent_record_id
        ):
            return record
    raise ValueError("Fattura di origine del riprocessamento non trovata")


def _compose_source_page_numbers(
    local_page_numbers: list[int],
    source_page_numbers: list[int],
) -> list[int]:
    if not source_page_numbers:
        return local_page_numbers
    mapped: list[int] = []
    for local_page_number in local_page_numbers:
        source_index = local_page_number - 1
        if source_index < 0 or source_index >= len(source_page_numbers):
            raise ValueError("Mappa pagine della fattura non valida")
        mapped.append(source_page_numbers[source_index])
    return mapped


def _page_numbers_payload(invoice: dict[str, Any]) -> list[int]:
    raw_page_numbers = invoice.get("pageNumbers")
    if not isinstance(raw_page_numbers, list):
        return []

    page_numbers: list[int] = []
    for page_number in raw_page_numbers:
        try:
            page_numbers.append(int(page_number))
        except (TypeError, ValueError):
            continue
    return page_numbers


def _invoice_export_row(record: dict[str, Any]) -> dict[str, str]:
    invoice = record.get("invoice") if isinstance(record.get("invoice"), dict) else {}
    field_values = _field_values(invoice)
    return {
        "saved_at": _date_part(str(record.get("saved_at") or "")),
        "nome_emittente": field_values.get("nome_emittente", ""),
        "codice_fiscale_emittente": field_values.get("codice_fiscale_emittente", ""),
        "p_iva": field_values.get("p_iva", ""),
        "condominio": field_values.get("condominio", ""),
        "codice_fiscale": field_values.get("codice_fiscale", ""),
        "numero_fattura": field_values.get("numero_fattura", ""),
        "data_fattura": field_values.get("data_fattura", ""),
        "totale_fattura": field_values.get("totale_fattura", ""),
        "ritenuta": field_values.get("ritenuta", ""),
        "totale_netto_fattura": field_values.get("totale_netto_fattura", ""),
        "pagata_il": field_values.get("pagata_il", ""),
        "source_file_name": str(record.get("source_file_name") or ""),
        "archived_pdf_path": str(record.get("archived_pdf_path") or ""),
        "archived_source_pdf_path": str(record.get("archived_source_pdf_path") or ""),
    }


def _field_values(invoice: dict[str, Any]) -> dict[str, str]:
    values: dict[str, str] = {}
    fields = invoice.get("fields", [])
    if not isinstance(fields, list):
        return values

    for field in fields:
        if not isinstance(field, dict):
            continue
        kind = str(field.get("kind") or "").strip()
        value = str(field.get("value") or "").strip()
        if kind:
            values[kind] = value

    for canonical_kind, aliases in INVOICE_FIELD_ALIASES.items():
        if values.get(canonical_kind):
            continue
        for alias in aliases:
            if values.get(alias):
                values[canonical_kind] = values[alias]
                break
    return values


def _date_part(value: str) -> str:
    return value[:10] if len(value) >= 10 else value


def _upsert_record(
    records: object,
    next_record: dict[str, Any],
) -> list[dict[str, Any]]:
    current_records = [
        record
        for record in records
        if isinstance(record, dict)
        and record.get("record_id") != next_record["record_id"]
    ]
    current_records.append(next_record)
    return current_records


def _required_string(payload: dict[str, Any], key: str) -> str:
    value = str(payload.get(key) or "").strip()
    if not value:
        raise ValueError(f"Campo obbligatorio mancante: {key}")
    return value


def _timestamp(now: datetime | None = None) -> str:
    return (now or datetime.now()).replace(microsecond=0).isoformat()


def _empty_registry() -> dict[str, Any]:
    return {
        "schema_version": INVOICE_REGISTRY_SCHEMA_VERSION,
        "updated_at": "",
        "records": [],
    }


def _write_registry(output_path: Path, payload: dict[str, Any]) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_name(f".{output_path.name}.writing")
    temporary_path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    temporary_path.replace(output_path)
