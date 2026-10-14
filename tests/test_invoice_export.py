import json
from io import BytesIO
from types import SimpleNamespace
from threading import Thread
from http.server import ThreadingHTTPServer
from urllib.request import Request, urlopen
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from estrattore.invoice_export import delete_export_history, export_invoices
from estrattore.invoice_registry_storage import load_invoice_registry
from estrattore.runtime_paths import get_runtime_paths
from scripts.desktop_app import DesktopPreviewApi
import scripts.desktop_app as desktop_app


def seed_export_runtime(root):
    paths = get_runtime_paths(root)
    paths.ensure_directories()
    records = []
    for index, cf in enumerate(["01234567890", "01234567890", "09876543210"], 1):
        values = {
            "condominio": "Condominio Via Roma / Scala A",
            "codice_fiscale": cf,
            "nome_emittente": "Impresa Rossi",
            "codice_fiscale_emittente": "00123456789",
            "p_iva": "00123456789",
            "numero_fattura": "0008",
            "data_fattura": "20-09-2026",
            "totale_fattura": "132.37",
            "ritenuta": "4.34",
            "totale_netto_fattura": "128.03",
            "pagata_il": "25-09-2026",
        }
        records.append({
            "record_id": f"r{index}", "status": "saved", "saved_at": "2026-09-28T14:00:00",
            "source_file_name": "sorgente.pdf", "archived_pdf_path": "archive/invoices/fattura.pdf",
            "invoice": {"fields": [{"kind": k, "value": v} for k, v in values.items()]},
        })
    paths.invoice_registry_file.write_text(json.dumps({"schema_version": 3, "records": records}), encoding="utf-8")
    registry_records = []
    for index, (cf, name) in enumerate([
        ("01234567890", "Condominio Via Roma / Scala A"),
        ("09876543210", "Condominio Via Roma / Scala A"),
        ("00123456789", "Impresa Rossi"),
    ]):
        registry_records.append({"registry_id": str(index), "source_row": index + 1,
                                 "codice_anagrafica": f"000{index}", "nominativo": name,
                                 "normalized_name": name.upper(), "codice_fiscale": cf,
                                 "tipologia": "Condominio" if index < 2 else "Altro soggetto"})
    paths.registry_cache_file.write_text(json.dumps({"schema_version": 1, "source_file": "test.xlsx", "records": registry_records}), encoding="utf-8")
    return paths


def artifact(paths, result):
    return paths.archive_dir / result["download_url"].removeprefix("/runtime/archive/")


def test_delivery_failure_does_not_create_export_history(tmp_path):
    paths = seed_export_runtime(tmp_path)
    original = paths.invoice_registry_file.read_bytes()

    def fail_delivery(_body, _name):
        raise OSError("destination unavailable")

    with pytest.raises(OSError, match="destination unavailable"):
        export_invoices(["r1"], "workbook", paths, deliver=fail_delivery)

    assert paths.invoice_registry_file.read_bytes() == original
    assert not (paths.archive_dir / "exports").exists()


def test_native_export_cancel_rename_and_redownload(tmp_path):
    paths = seed_export_runtime(tmp_path / "app")
    chosen = [None, tmp_path / "Custom invoice name.xlsx", tmp_path / "Second copy.xlsx"]
    dialog_calls = []

    def dialog(kind, **kwargs):
        dialog_calls.append((kind, kwargs))
        destination = chosen.pop(0)
        return (str(destination),) if destination else None

    bridge = DesktopPreviewApi(SimpleNamespace(FileDialog=SimpleNamespace(SAVE=30)), 8765, paths.archive_dir, paths)
    bridge.bind_window(SimpleNamespace(create_file_dialog=dialog))
    original = paths.invoice_registry_file.read_bytes()
    assert bridge.export_invoices_native(["r1"], "workbook")["cancelled"] is True
    assert paths.invoice_registry_file.read_bytes() == original
    assert not (paths.archive_dir / "exports").exists()

    result = bridge.export_invoices_native(["r1"], "workbook")
    assert result["ok"] is True
    workbook = load_workbook(tmp_path / "Custom invoice name.xlsx")
    assert workbook.active["D5"].value == "0008"
    workbook.close()
    assert len(load_invoice_registry(paths)["export_history"]) == 1
    assert bridge.save_existing_export(result["export_id"])["ok"] is True
    assert (tmp_path / "Second copy.xlsx").read_bytes() == (tmp_path / "Custom invoice name.xlsx").read_bytes()
    assert len(load_invoice_registry(paths)["export_history"]) == 1
    assert all(kind == 30 for kind, _ in dialog_calls)


def test_native_report_save_does_not_allow_wrong_extension(tmp_path):
    paths = get_runtime_paths(tmp_path / "app")
    paths.ensure_directories()
    paths.registry_duplicate_report_xlsx_file.write_bytes(b"report")
    bridge = DesktopPreviewApi(SimpleNamespace(FileDialog=SimpleNamespace(SAVE=30)), 8765, paths.archive_dir, paths)
    bridge.bind_window(SimpleNamespace(create_file_dialog=lambda *_args, **_kwargs: (str(tmp_path / "wrong.zip"),)))
    with pytest.raises(ValueError, match="estensione"):
        bridge.save_named_export("cleaning_report")
    assert not (tmp_path / "wrong.zip").exists()


def test_native_named_exports_use_selected_names(tmp_path):
    paths = seed_export_runtime(tmp_path / "app")
    paths.registry_duplicate_report_xlsx_file.write_bytes(b"report")
    filenames = iter(("renamed-registry.xlsx", "renamed-report.xlsx", "renamed-invoices.xlsx"))
    bridge = DesktopPreviewApi(SimpleNamespace(FileDialog=SimpleNamespace(SAVE=30)), 8765, paths.archive_dir, paths)
    bridge.bind_window(SimpleNamespace(create_file_dialog=lambda *_args, **_kwargs: (str(tmp_path / next(filenames)),)))

    assert bridge.save_named_export("registry")["ok"]
    assert load_workbook(tmp_path / "renamed-registry.xlsx").active.max_row >= 2
    assert bridge.save_named_export("cleaning_report")["ok"]
    assert (tmp_path / "renamed-report.xlsx").read_bytes() == b"report"
    assert bridge.save_named_export("invoice_registry")["ok"]
    assert load_workbook(tmp_path / "renamed-invoices.xlsx").active.max_row >= 2


def test_native_invoice_export_uses_zip_for_summary(tmp_path):
    paths = seed_export_runtime(tmp_path / "app")
    chosen = tmp_path / "riepilogo personalizzato.zip"
    suggestions = []

    def dialog(_kind, **kwargs):
        suggestions.append(kwargs["save_filename"])
        return (str(chosen),)

    bridge = DesktopPreviewApi(SimpleNamespace(FileDialog=SimpleNamespace(SAVE=30)), 8765, paths.archive_dir, paths)
    bridge.bind_window(SimpleNamespace(create_file_dialog=dialog))
    result = bridge.export_invoices_native(["r1"], "workbook", include_summary=True)
    assert result["ok"] is True
    assert suggestions == ["fatture_condomini.zip"]
    with ZipFile(chosen) as archive:
        assert sorted(archive.namelist()) == ["fatture_condomini.xlsx", "riepilogo_ritenute.xlsx"]


def test_native_atomic_save_failure_preserves_existing_file(tmp_path, monkeypatch):
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"old")

    def fail_replace(_source, _target):
        raise PermissionError("file locked")

    monkeypatch.setattr(desktop_app.os, "replace", fail_replace)
    with pytest.raises(PermissionError, match="file locked"):
        DesktopPreviewApi._write_export_file(target, b"new")
    assert target.read_bytes() == b"old"
    assert list(tmp_path.iterdir()) == [target]


def test_grouped_workbook_and_delivery_metadata(tmp_path):
    paths = seed_export_runtime(tmp_path)
    result = export_invoices(["r1", "r2", "r3"], "workbook", paths)
    workbook = load_workbook(artifact(paths, result))
    assert len(workbook.worksheets) == 2
    assert len(set(name.casefold() for name in workbook.sheetnames)) == 2
    sheet = workbook.worksheets[0]
    assert [sheet.cell(2, i).value for i in range(1, 5)] == [
        "Cod. anagrafica", "0000", "C.F. condominio", "01234567890",
    ]
    assert [sheet.cell(4, i).value for i in range(1, 6)] == ["Cod. cliente", "Denominazione", "Data doc.", "N. doc", "Compenso"]
    assert sheet["A5"].value == "0002"
    assert sheet["D5"].value == "0008"
    assert sheet["E5"].value == 132.37
    assert sheet["H5"].value == 4.34
    assert sheet["C5"].value.day == 20
    assert sheet["E5"].fill.fgColor.rgb == "00FFFFFF"
    assert sheet["F5"].fill.fgColor.rgb == "00F7F8F7"
    assert sheet["A1"].fill.fgColor.rgb == "00E7F1EA"
    assert sheet["A4"].fill.fgColor.rgb == "00E0EDE4"
    assert sheet.max_row == 6
    assert sheet.max_column == 10
    assert sheet["J4"].value == "Pagata il"
    assert sheet["J5"].value.day == 25
    assert sheet["K4"].value is None
    for coordinate in ("A1", "A2", "B2", "C2", "D2", "A4", "B5", "E5", "J5"):
        assert sheet[coordinate].alignment.horizontal == "center"
    assert sheet.print_area.endswith("$A$1:$J$6")
    assert sheet.freeze_panes is None
    assert sheet.auto_filter.ref is None
    saved = load_invoice_registry(paths)["records"]
    assert all(r["status"] == "saved" and r["last_exported_at"] for r in saved)
    assert all(r["last_export_url"] == result["download_url"] for r in saved)


def test_separate_export_subset_and_reexport_preserve_previous_download(tmp_path):
    paths = seed_export_runtime(tmp_path)
    first = export_invoices(["r1", "r3"], "separate", paths)
    with ZipFile(artifact(paths, first)) as archive:
        assert len(archive.namelist()) == 2
        for name in archive.namelist():
            workbook = load_workbook(BytesIO(archive.read(name)))
            assert len(workbook.worksheets) == 1
            assert workbook.active.freeze_panes is None
    second = export_invoices(["r1"], "workbook", paths)
    records = {r["record_id"]: r for r in load_invoice_registry(paths)["records"]}
    assert "last_exported_at" not in records["r2"]
    assert records["r3"]["last_export_url"] == first["download_url"]
    assert records["r1"]["last_exported_at"] == second["exported_at"]
    assert artifact(paths, first).exists()


def test_failures_do_not_mark_rows_exported(tmp_path, monkeypatch):
    from estrattore import invoice_export
    paths = seed_export_runtime(tmp_path)
    before = paths.invoice_registry_file.read_bytes()
    with pytest.raises(ValueError):
        export_invoices(["r1", "missing"], "workbook", paths)
    def fail(*args):
        raise OSError("disk full")
    monkeypatch.setattr(invoice_export, "_write_registry", fail)
    with pytest.raises(OSError):
        export_invoices(["r1"], "workbook", paths)
    assert paths.invoice_registry_file.read_bytes() == before
    assert not list((paths.archive_dir / "exports").glob("*/*"))


def test_generation_failure_keeps_pending_rows_and_creates_no_artifact(tmp_path, monkeypatch):
    from estrattore import invoice_export
    paths = seed_export_runtime(tmp_path)
    before = paths.invoice_registry_file.read_bytes()
    def fail(*args):
        raise OSError("generation failed")
    monkeypatch.setattr(invoice_export, "_workbook", fail)
    with pytest.raises(OSError):
        export_invoices(["r1"], "workbook", paths)
    assert paths.invoice_registry_file.read_bytes() == before
    assert not (paths.archive_dir / "exports").exists()


def test_selected_registry_row_disambiguates_client_code(tmp_path):
    paths = seed_export_runtime(tmp_path)
    cache = json.loads(paths.registry_cache_file.read_text())
    cache["records"].append({**cache["records"][-1], "registry_id": "other", "source_row": 99, "codice_anagrafica": "9999"})
    paths.registry_cache_file.write_text(json.dumps(cache))
    payload = json.loads(paths.invoice_registry_file.read_text())
    field = next(f for f in payload["records"][0]["invoice"]["fields"] if f["kind"] == "nome_emittente")
    field.update(registry_id="other", registry_source_row=99)
    paths.invoice_registry_file.write_text(json.dumps(payload))
    result = export_invoices(["r1"], "workbook", paths)
    assert load_workbook(artifact(paths, result)).active["A5"].value == "9999"


def test_text_is_not_formula_and_ambiguous_code_is_blank(tmp_path):
    paths = seed_export_runtime(tmp_path)
    payload = json.loads(paths.invoice_registry_file.read_text())
    fields = payload["records"][0]["invoice"]["fields"]
    next(f for f in fields if f["kind"] == "numero_fattura")["value"] = '=HYPERLINK("x")'
    paths.invoice_registry_file.write_text(json.dumps(payload))
    cache = json.loads(paths.registry_cache_file.read_text())
    cache["records"].append({**cache["records"][-1], "registry_id": "other", "source_row": 99, "codice_anagrafica": "9999"})
    paths.registry_cache_file.write_text(json.dumps(cache))
    result = export_invoices(["r1"], "workbook", paths)
    sheet = load_workbook(artifact(paths, result)).active
    assert sheet["A5"].value is None
    assert sheet["D5"].data_type == "s"
    assert sheet["D5"].value.startswith("=")


def test_history_membership_survives_overlapping_reexports(tmp_path):
    paths = seed_export_runtime(tmp_path)
    first = export_invoices(["r1", "r2"], "separate", paths)
    original = load_invoice_registry(paths)["export_history"][0]
    export_invoices(["r1"], "workbook", paths)
    replay = export_invoices(["r3"], "workbook", paths, history_id=first["export_id"])
    assert replay["invoice_count"] == 2  # Membership comes from the receipt, not submitted IDs.
    history = load_invoice_registry(paths)["export_history"]
    assert len(history) == 3
    assert history[0] == original
    assert history[-1]["source_export_id"] == first["export_id"]
    assert {item["record_id"] for item in history[-1]["invoices"]} == {"r1", "r2"}


@pytest.mark.parametrize("removal", ["missing", "hidden", "discarded"])
def test_history_reexport_rejects_unavailable_member_atomically(tmp_path, removal):
    paths = seed_export_runtime(tmp_path)
    exported = export_invoices(["r1", "r2"], "workbook", paths)
    registry = load_invoice_registry(paths)
    if removal == "missing":
        registry["records"] = [r for r in registry["records"] if r["record_id"] != "r2"]
    elif removal == "hidden":
        registry["records"][1]["registry_visible"] = False
    else:
        registry["records"][1]["status"] = "discarded"
    paths.invoice_registry_file.write_text(json.dumps(registry))
    before = paths.invoice_registry_file.read_bytes()
    files = set((paths.archive_dir / "exports").rglob("*"))
    with pytest.raises(ValueError, match="Non e possibile riesportare"):
        export_invoices(None, "separate", paths, history_id=exported["export_id"])
    assert paths.invoice_registry_file.read_bytes() == before
    assert set((paths.archive_dir / "exports").rglob("*")) == files
    assert len(load_invoice_registry(paths)["export_history"][0]["invoices"]) == 2


def test_history_delete_preserves_invoices_and_cleans_only_unreferenced_artifact(tmp_path):
    paths = seed_export_runtime(tmp_path)
    first = export_invoices(["r1"], "workbook", paths)
    second = export_invoices(["r1"], "separate", paths)
    before = load_invoice_registry(paths)["records"]
    delete_export_history(first["export_id"], paths)
    assert not artifact(paths, first).exists()
    assert artifact(paths, second).exists()
    delete_export_history(second["export_id"], paths)
    assert artifact(paths, second).exists()  # Still the invoice's latest download.
    assert load_invoice_registry(paths)["records"] == before
    assert load_invoice_registry(paths)["export_history"] == []
    with pytest.raises(ValueError, match="non trovata"):
        export_invoices(None, "workbook", paths, history_id=first["export_id"])


@pytest.mark.parametrize("handler_name", ["local_app", "registry_api"])
def test_real_local_app_export_and_download(tmp_path, monkeypatch, handler_name):
    from scripts import registry_api, local_app
    paths = seed_export_runtime(tmp_path)
    monkeypatch.setattr(registry_api, "RUNTIME_PATHS", paths)
    monkeypatch.setattr(local_app, "RUNTIME_PATHS", paths)
    handler = local_app.LocalAppHandler if handler_name == "local_app" else registry_api.RegistryApiHandler
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        request = Request(base + "/invoices/export", data=json.dumps({"record_ids": ["r1"], "mode": "workbook"}).encode(), headers={"Content-Type": "application/json"})
        with urlopen(request) as response:
            result = json.load(response)
        with urlopen(base + result["download_url"]) as response:
            workbook = load_workbook(BytesIO(response.read()))
            assert workbook.active["E5"].value == 132.37
        with urlopen(base + "/invoices/registry") as response:
            payload = json.load(response)["registry"]
            records = payload["records"]
            assert records[0]["last_exported_at"] == result["exported_at"]
            assert payload["export_history"][0]["export_id"] == result["export_id"]
        request = Request(base + "/invoices/export", data=json.dumps({"history_id": result["export_id"], "mode": "separate", "include_summary": True}).encode(), headers={"Content-Type": "application/json"})
        with urlopen(request) as response:
            replay = json.load(response)
            assert replay["invoice_count"] == 1
        with urlopen(base + replay["download_url"]) as response:
            with ZipFile(BytesIO(response.read())) as bundle:
                summary = load_workbook(BytesIO(bundle.read("riepilogo_ritenute.xlsx")))
                assert summary.active["E3"].value == 4.34
        request = Request(base + "/invoices/export-history/delete", data=json.dumps({"export_id": result["export_id"]}).encode(), headers={"Content-Type": "application/json"})
        with urlopen(request) as response:
            assert json.load(response)["ok"]
        with urlopen(base + "/invoices/registry") as response:
            payload = json.load(response)["registry"]
            assert len(payload["export_history"]) == 1
            assert len(payload["records"]) == 3
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


@pytest.mark.parametrize("mode, files", [("workbook", 2), ("separate", 3)])
def test_optional_summary_has_exact_membership_and_totals(tmp_path, mode, files):
    paths = seed_export_runtime(tmp_path)
    result = export_invoices(["r1", "r2", "r3"], mode, paths, include_summary=True)
    with ZipFile(artifact(paths, result)) as bundle:
        assert len(bundle.namelist()) == files
        sheet = load_workbook(BytesIO(bundle.read("riepilogo_ritenute.xlsx"))).active
        assert sheet["A2"].value == "Cod. anagrafica"
        assert sheet["E2"].value == "Ritenute"
        assert sheet["A3"].value == "0000"
        assert sheet["C3"].value == "01234567890"
        assert sheet["D3"].value == 2
        assert sheet["E3"].value == 8.68
        assert sheet["E4"].value == 4.34
        assert sheet["B5"].value == "Totale ritenute"
        assert sheet["D5"].value == 3
        assert sheet["E5"].value == 13.02
        assert sheet.max_column == 5
        assert sheet["A1"].value == "Riepilogo ritenute delle fatture esportate"
        assert sheet.freeze_panes is None
    receipt = load_invoice_registry(paths)["export_history"][0]
    assert receipt["include_summary"] is True
    replay = export_invoices(None, "workbook", paths, history_id=result["export_id"], include_summary=False)
    assert artifact(paths, replay).suffix == ".xlsx"
    assert len(load_workbook(artifact(paths, replay)).worksheets) == 2


@pytest.mark.parametrize("value", ["", "non valido", "-4.34"])
def test_summary_flags_unknown_amount_without_assuming_zero(tmp_path, value):
    paths = seed_export_runtime(tmp_path)
    registry = load_invoice_registry(paths)
    for record in registry["records"]:
        next(f for f in record["invoice"]["fields"] if f["kind"] == "ritenuta")["value"] = value
    paths.invoice_registry_file.write_text(json.dumps(registry))
    result = export_invoices(["r1"], "workbook", paths, include_summary=True)
    with ZipFile(artifact(paths, result)) as bundle:
        sheet = load_workbook(BytesIO(bundle.read("riepilogo_ritenute.xlsx"))).active
        assert sheet["E3"].value is None
        assert sheet["E4"].value is None
        assert sheet["F2"].value is None


def test_summary_option_rejects_non_boolean_without_changes(tmp_path):
    paths = seed_export_runtime(tmp_path)
    before = paths.invoice_registry_file.read_bytes()
    with pytest.raises(ValueError, match="riepilogo"):
        export_invoices(["r1"], "workbook", paths, include_summary="false")
    assert paths.invoice_registry_file.read_bytes() == before
