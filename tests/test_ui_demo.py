from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from estrattore.invoice_registry_storage import load_invoice_registry
from estrattore.processing_storage import list_review_batches
from estrattore.registry import load_registry_cache
from estrattore.runtime_paths import get_runtime_paths
from estrattore.source_archive_storage import list_archived_sources
from scripts.desktop_app import APP_MUTEX_NAME, _mutex_name_for_paths
from scripts.ui_demo import create_demo_data


def test_ui_demo_populates_real_isolated_stores(tmp_path: Path) -> None:
    root = tmp_path / "ui-demo"
    assert create_demo_data(root) == root
    paths = get_runtime_paths(root)

    assert len(load_registry_cache(paths.registry_cache_file).records) == 2
    workbook = load_workbook(paths.registry_input_file, read_only=True)
    assert workbook.active.max_row == 3
    workbook.close()

    reviews = list_review_batches(paths)
    assert len(reviews) == 1
    assert reviews[0].invoice_count == 1

    registry = load_invoice_registry(paths)
    records = registry["records"]
    assert [record["status"] for record in records] == ["saved", "saved", "discarded"]
    assert records[0].get("last_exported_at") is None
    assert records[1].get("last_exported_at")
    assert len(registry["export_history"]) == 1
    for record in records:
        pdf = root / record["archived_pdf_path"]
        assert pdf.is_file()
        assert len(PdfReader(pdf).pages) == 1
        assert (root / record["archived_source_pdf_path"]).is_file()

    sources = list_archived_sources(paths)
    assert len(sources) == 1
    assert len(PdfReader(root / sources[0].relative_path).pages) == 4

    assert create_demo_data(root) == root
    assert len(load_invoice_registry(paths)["records"]) == 3
    assert _mutex_name_for_paths(paths, False) == APP_MUTEX_NAME
    assert _mutex_name_for_paths(paths, True) != APP_MUTEX_NAME
    with pytest.raises(ValueError, match="dati dimostrativi"):
        _mutex_name_for_paths(get_runtime_paths(tmp_path / "other-data"), True)
