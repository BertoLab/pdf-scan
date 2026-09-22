import json
import sqlite3
from contextlib import closing
from datetime import datetime
from pathlib import Path
import shutil
import uuid

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

from estrattore.batch_storage import PdfBatchSource, import_pdf_batch
from estrattore.invoice_pdf_archive import ArchivedInvoicePdf
from estrattore.invoice_registry_storage import (
    delete_invoice_archive,
    discard_reviewed_invoice,
    load_invoice_registry,
    purge_discarded_invoice_records,
    save_invoice_registry_excel,
    save_invoice_registry_records,
    save_reviewed_invoice,
)
from estrattore.processing_storage import (
    list_review_batches,
    process_imported_batch,
    review_batch_dir_from_base_path,
)
from estrattore.runtime_paths import get_runtime_paths
from estrattore.source_archive_storage import promote_batch_sources


def _workspace() -> Path:
    root = Path("invoice-registry-storage-test-manual") / uuid.uuid4().hex
    root.mkdir(parents=True)
    return root


def _clean_workspace(path: Path) -> None:
    shutil.rmtree(path, ignore_errors=True)


def test_save_reviewed_invoice_persists_snapshot_and_removes_it_from_review(
    monkeypatch,
) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        invoice_payload = _invoice_payload("invoice-001")
        _mock_pdf_archive(monkeypatch, paths)

        record = save_reviewed_invoice(
            {
                "selected_batch_id": review_batch.id,
                "review_batch_base_path": review_batch.base_path,
                "invoice": invoice_payload,
            },
            runtime_paths=paths,
            now=datetime(2026, 8, 10, 12, 0, 0),
        )
        repeated_record = save_reviewed_invoice(
            {
                "selected_batch_id": review_batch.id,
                "review_batch_base_path": review_batch.base_path,
                "invoice": invoice_payload,
            },
            runtime_paths=paths,
            now=datetime(2026, 8, 10, 12, 1, 0),
        )

        registry = load_invoice_registry(paths)
        assert repeated_record == record
        assert record["status"] == "saved"
        assert record["archived_pdf_path"] == "archive/invoices/fake_invoice.pdf"
        assert record["archived_pdf_url"] == "/runtime/archive/invoices/fake_invoice.pdf"
        assert record["archived_pdf_page_numbers"] == [1]
        assert record["archived_source_pdf_path"] == "archive/sources/source.pdf"
        assert record["archived_source_pdf_url"] == "/runtime/archive/sources/source.pdf"
        assert registry["records"][0]["invoice"]["fields"][0]["value"] == "85,79"
        assert list_review_batches(paths) == []

        assert not review_batch_dir_from_base_path(
            review_batch.base_path,
            paths,
            must_exist=False,
        ).exists()
        assert not (
            paths.manifests_dir
            / "process_20260810_110000"
            / "documents"
            / "doc_000001"
        ).exists()
        assert (
            paths.batches_dir
            / "batch_20260810_100000"
            / "input"
            / "doc_000001.pdf"
        ).exists()
    finally:
        _clean_workspace(workspace)


def test_save_reviewed_invoice_recovers_feedback_commit_after_registry_write(
    monkeypatch,
) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        payload = {
            "selected_batch_id": review_batch.id,
            "review_batch_base_path": review_batch.base_path,
            "invoice": _invoice_payload("invoice-001"),
        }
        _mock_pdf_archive(monkeypatch, paths)
        from estrattore import invoice_registry_storage

        real_commit = invoice_registry_storage.commit_review_operation
        commit_attempts = 0

        def fail_first_commit(operation_id, runtime_paths):
            nonlocal commit_attempts
            commit_attempts += 1
            if commit_attempts == 1:
                raise RuntimeError("simulated feedback commit interruption")
            return real_commit(operation_id, runtime_paths)

        monkeypatch.setattr(
            invoice_registry_storage,
            "commit_review_operation",
            fail_first_commit,
        )

        saved_record = save_reviewed_invoice(payload, runtime_paths=paths)

        assert len(load_invoice_registry(paths)["records"]) == 1
        assert len(list_review_batches(paths)) == 0

        changed_payload = {
            **payload,
            "invoice": _record_with_field_value(
                {"invoice": payload["invoice"]}, "totale_fattura", "90,00"
            )["invoice"],
        }
        repeated_record = save_reviewed_invoice(changed_payload, runtime_paths=paths)
        assert repeated_record["invoice"]["fields"][0]["value"] == "85,79"

        with closing(sqlite3.connect(paths.feedback_database_file)) as connection:
            revision_count = connection.execute(
                "SELECT COUNT(*) FROM review_revisions"
            ).fetchone()[0]
            pending_statuses = connection.execute(
                "SELECT status FROM pending_operations"
            ).fetchall()
        assert revision_count == 1
        assert pending_statuses == [("completed",)]
    finally:
        _clean_workspace(workspace)


def test_collection_prepare_failure_does_not_block_invoice_save(monkeypatch) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        _mock_pdf_archive(monkeypatch, paths)
        monkeypatch.setattr(
            "estrattore.invoice_registry_storage.prepare_review_operation",
            lambda **_kwargs: (_ for _ in ()).throw(OSError("collection unavailable")),
        )

        record = save_reviewed_invoice(
            {
                "review_batch_base_path": review_batch.base_path,
                "invoice": _invoice_payload("invoice-001"),
            },
            runtime_paths=paths,
        )

        assert record["status"] == "saved"
        assert load_invoice_registry(paths)["records"][0]["record_id"] == record["record_id"]
        assert list_review_batches(paths) == []
        with closing(sqlite3.connect(paths.feedback_database_file)) as connection:
            issue = connection.execute(
                "SELECT stage, error_message FROM capture_issues WHERE resolved_at = ''"
            ).fetchone()
        assert issue == ("invoice_save_prepare", "collection unavailable")
    finally:
        _clean_workspace(workspace)


def test_invoice_persistence_accepts_three_digit_withholding_amount(monkeypatch) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        invoice_payload = _invoice_payload("invoice-001")
        invoice_payload["fields"].append(
            {"kind": "ritenuta", "label": "Ritenuta", "value": "100"}
        )
        _mock_pdf_archive(monkeypatch, paths)

        record = save_reviewed_invoice(
            {
                "selected_batch_id": review_batch.id,
                "review_batch_base_path": review_batch.base_path,
                "invoice": invoice_payload,
            },
            runtime_paths=paths,
        )

        assert record["status"] == "saved"
    finally:
        _clean_workspace(workspace)


def test_invoice_persistence_rejects_withholding_percentage(monkeypatch) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        invoice_payload = _invoice_payload("invoice-001")
        invoice_payload["fields"].append(
            {"kind": "ritenuta", "label": "Ritenuta", "value": "4%"}
        )
        _mock_pdf_archive(monkeypatch, paths)

        with pytest.raises(ValueError, match="Importo non valido"):
            save_reviewed_invoice(
                {
                    "selected_batch_id": review_batch.id,
                    "review_batch_base_path": review_batch.base_path,
                    "invoice": invoice_payload,
                },
                runtime_paths=paths,
            )
    finally:
        _clean_workspace(workspace)


@pytest.mark.parametrize(
    ("field_kind", "label", "invalid_value"),
    [
        ("totale_fattura", "Totale fattura", "cento"),
        ("ritenuta", "Ritenuta", "+4.00"),
        ("totale_netto_fattura", "Totale netto fattura", "-96.00"),
    ],
)
def test_invoice_persistence_rejects_invalid_accounting_amounts(
    monkeypatch,
    field_kind: str,
    label: str,
    invalid_value: str,
) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        invoice_payload = _invoice_payload("invoice-001")
        invoice_payload["fields"].append(
            {"kind": field_kind, "label": label, "value": invalid_value}
        )
        _mock_pdf_archive(monkeypatch, paths)

        with pytest.raises(ValueError, match="Importo non valido"):
            save_reviewed_invoice(
                {
                    "selected_batch_id": review_batch.id,
                    "review_batch_base_path": review_batch.base_path,
                    "invoice": invoice_payload,
                },
                runtime_paths=paths,
            )
    finally:
        _clean_workspace(workspace)


def test_invoice_registry_edit_rejects_inconsistent_withholding_amount(monkeypatch) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        invoice_payload = _invoice_payload("invoice-001")
        invoice_payload["fields"].extend(
            [
                {"kind": "totale_fattura", "label": "Totale fattura", "value": "100.00"},
                {"kind": "ritenuta", "label": "Ritenuta", "value": "10.00"},
                {
                    "kind": "totale_netto_fattura",
                    "label": "Totale netto fattura",
                    "value": "90.00",
                },
            ]
        )
        _mock_pdf_archive(monkeypatch, paths)
        record = save_reviewed_invoice(
            {
                "selected_batch_id": review_batch.id,
                "review_batch_base_path": review_batch.base_path,
                "invoice": invoice_payload,
            },
            runtime_paths=paths,
        )
        edited_record = {
            **record,
            "invoice": {
                **record["invoice"],
                "fields": [
                    {
                        **field,
                        "value": "20.00" if field["kind"] == "ritenuta" else field["value"],
                    }
                    for field in record["invoice"]["fields"]
                ],
            },
        }

        with pytest.raises(ValueError, match="Valori contabili non coerenti"):
            save_invoice_registry_records([edited_record], runtime_paths=paths)
    finally:
        _clean_workspace(workspace)


def test_invoice_registry_edit_recovers_feedback_after_business_write(monkeypatch) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        _mock_pdf_archive(monkeypatch, paths)
        record = save_reviewed_invoice(
            {
                "selected_batch_id": review_batch.id,
                "review_batch_base_path": review_batch.base_path,
                "invoice": _invoice_payload("invoice-001"),
            },
            runtime_paths=paths,
        )
        edited_record = _record_with_field_value(record, "totale_fattura", "90,00")
        from estrattore import invoice_registry_storage

        real_commit = invoice_registry_storage.commit_review_operation
        commit_attempts = 0

        def fail_first_commit(operation_id, runtime_paths):
            nonlocal commit_attempts
            commit_attempts += 1
            if commit_attempts == 1:
                raise RuntimeError("simulated registry feedback interruption")
            return real_commit(operation_id, runtime_paths)

        monkeypatch.setattr(
            invoice_registry_storage,
            "commit_review_operation",
            fail_first_commit,
        )

        save_invoice_registry_records([edited_record], runtime_paths=paths)

        persisted = load_invoice_registry(paths)["records"][0]
        assert _invoice_field_values_for_test(persisted)["totale_fattura"] == "90,00"
        with closing(sqlite3.connect(paths.feedback_database_file)) as connection:
            assert connection.execute(
                "SELECT COUNT(*) FROM review_revisions"
            ).fetchone()[0] == 2
            assert connection.execute(
                "SELECT status FROM pending_operations WHERE event_kind = 'registry_edit'"
            ).fetchone()[0] == "completed"

        save_invoice_registry_records([edited_record], runtime_paths=paths)

        with closing(sqlite3.connect(paths.feedback_database_file)) as connection:
            assert connection.execute(
                "SELECT COUNT(*) FROM review_revisions"
            ).fetchone()[0] == 2
            assert connection.execute(
                "SELECT status FROM pending_operations WHERE event_kind = 'registry_edit'"
            ).fetchone()[0] == "completed"
    finally:
        _clean_workspace(workspace)


def test_collection_prepare_failure_does_not_block_registry_edit(monkeypatch) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        _mock_pdf_archive(monkeypatch, paths)
        record = save_reviewed_invoice(
            {
                "review_batch_base_path": review_batch.base_path,
                "invoice": _invoice_payload("invoice-001"),
            },
            runtime_paths=paths,
        )
        changed = _record_with_field_value(record, "totale_fattura", "90,00")
        monkeypatch.setattr(
            "estrattore.invoice_registry_storage.prepare_registry_edit_operation",
            lambda **_kwargs: (_ for _ in ()).throw(OSError("collection unavailable")),
        )

        registry = save_invoice_registry_records([changed], runtime_paths=paths)

        assert _invoice_field_values_for_test(registry["records"][0])["totale_fattura"] == "90,00"
        with closing(sqlite3.connect(paths.feedback_database_file)) as connection:
            issue = connection.execute(
                """
                SELECT stage, error_message FROM capture_issues
                WHERE stage = 'registry_edit_prepare' AND resolved_at = ''
                """
            ).fetchone()
        assert issue == ("registry_edit_prepare", "collection unavailable")
    finally:
        _clean_workspace(workspace)


def test_registry_edit_returning_to_an_old_value_creates_a_new_revision(
    monkeypatch,
) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        _mock_pdf_archive(monkeypatch, paths)
        record = save_reviewed_invoice(
            {
                "selected_batch_id": review_batch.id,
                "review_batch_base_path": review_batch.base_path,
                "invoice": _invoice_payload("invoice-001"),
            },
            runtime_paths=paths,
        )
        changed = _record_with_field_value(record, "totale_fattura", "90,00")
        save_invoice_registry_records([changed], runtime_paths=paths)
        restored = _record_with_field_value(changed, "totale_fattura", "85,79")
        save_invoice_registry_records([restored], runtime_paths=paths)

        with closing(sqlite3.connect(paths.feedback_database_file)) as connection:
            revisions = connection.execute(
                """
                SELECT event_kind, accepted_fields_json
                FROM review_revisions ORDER BY rowid
                """
            ).fetchall()
        assert [revision[0] for revision in revisions] == [
            "invoice_save",
            "registry_edit",
            "registry_edit",
        ]
        assert json.loads(revisions[-1][1])[0]["value"] == "85,79"
    finally:
        _clean_workspace(workspace)


def test_registry_edit_repeated_transition_creates_a_revision_per_occurrence(
    monkeypatch,
) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        _mock_pdf_archive(monkeypatch, paths)
        record = save_reviewed_invoice(
            {
                "review_batch_base_path": review_batch.base_path,
                "invoice": _invoice_payload("invoice-001"),
            },
            runtime_paths=paths,
        )
        for value in ("90,00", "85,79", "90,00"):
            record = _record_with_field_value(record, "totale_fattura", value)
            registry = save_invoice_registry_records([record], runtime_paths=paths)
            record = registry["records"][0]

        with closing(sqlite3.connect(paths.feedback_database_file)) as connection:
            revisions = connection.execute(
                "SELECT accepted_fields_json FROM review_revisions ORDER BY rowid"
            ).fetchall()
        assert len(revisions) == 4
        assert json.loads(revisions[-1][0])[0]["value"] == "90,00"
    finally:
        _clean_workspace(workspace)


def test_registry_edit_replaces_stale_registry_provenance_with_manual_action(
    monkeypatch,
) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        _mock_pdf_archive(monkeypatch, paths)
        invoice = _invoice_payload("invoice-001")
        invoice["fields"][0].update(
            review_interaction="registry_selection",
            review_value_origin="explicit_registry",
            review_action_id="old-registry-action",
            registry_id="registry-1",
            review_registry_snapshot={"nominativo": "ACME"},
        )
        record = save_reviewed_invoice(
            {"review_batch_base_path": review_batch.base_path, "invoice": invoice},
            runtime_paths=paths,
        )
        changed = _record_with_field_value(record, "totale_fattura", "90,00")
        save_invoice_registry_records([changed], runtime_paths=paths)

        with closing(sqlite3.connect(paths.feedback_database_file)) as connection:
            provenance_json = connection.execute(
                """
                SELECT field_provenance_json FROM review_revisions
                WHERE event_kind = 'registry_edit'
                """
            ).fetchone()[0]
        provenance = json.loads(provenance_json)[0]
        assert provenance["interaction"] == "manual_input"
        assert provenance["value_origin"] == "manual_input"
        assert provenance["action_id"] != "old-registry-action"
        assert provenance["registry_id"] is None
        assert provenance["registry_snapshot"] is None
    finally:
        _clean_workspace(workspace)


def test_save_reviewed_invoice_accepts_composite_frontend_invoice_id(monkeypatch) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        invoice_payload = _invoice_payload("doc_000001:invoice-001")
        invoice_payload["manifestInvoiceId"] = "invoice-001"
        _mock_pdf_archive(monkeypatch, paths)

        save_reviewed_invoice(
            {
                "selected_batch_id": review_batch.id,
                "review_batch_base_path": review_batch.base_path,
                "invoice": invoice_payload,
            },
            runtime_paths=paths,
            now=datetime(2026, 8, 10, 12, 0, 0),
        )

        assert list_review_batches(paths) == []
    finally:
        _clean_workspace(workspace)


def test_discard_reviewed_invoice_archives_history_once_and_closes_review(
    monkeypatch,
) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        payload = {
            "selected_batch_id": review_batch.id,
            "review_batch_base_path": review_batch.base_path,
            "invoice": _invoice_payload("invoice-001"),
        }
        _mock_pdf_archive(monkeypatch, paths)

        discard_reviewed_invoice(
            payload,
            runtime_paths=paths,
            now=datetime(2026, 8, 10, 12, 0, 0),
        )
        discard_reviewed_invoice(
            payload,
            runtime_paths=paths,
            now=datetime(2026, 8, 10, 12, 1, 0),
        )

        registry = load_invoice_registry(paths)
        assert len(registry["records"]) == 1
        discarded = registry["records"][0]
        assert discarded["status"] == "discarded"
        assert discarded["discarded_at"] == "2026-08-10T12:00:00"
        assert discarded["display_name"] == "source - pagina 1"
        assert discarded["archived_pdf_path"] == "archive/invoices/fake_invoice.pdf"
        assert discarded["invoice"]["invoiceId"] == "invoice-001"
        assert list_review_batches(paths) == []
    finally:
        _clean_workspace(workspace)


def test_reprocessed_discarded_invoice_preserves_original_page_lineage() -> None:
    workspace = _workspace()
    try:
        paths = get_runtime_paths(workspace / "app-data")
        source_pdf = workspace / "documento.pdf"
        _write_blank_pdf(source_pdf, 3)
        source_batch = import_pdf_batch([source_pdf], runtime_paths=paths)

        def parent_processor(_source_path: Path, output_dir: Path, first_number: int):
            return _write_manifest(output_dir, first_number, [1, 3])

        source_run = process_imported_batch(
            source_batch.batch_id,
            runtime_paths=paths,
            document_processor=parent_processor,
        )
        promote_batch_sources(source_batch, source_run.process_id, paths)
        parent_invoice = _invoice_payload("invoice-001")
        parent_invoice["pageNumbers"] = [1, 3]
        discarded = discard_reviewed_invoice(
            {
                "selected_batch_id": source_run.review_batches[0].id,
                "review_batch_base_path": source_run.review_batches[0].base_path,
                "invoice": parent_invoice,
            },
            runtime_paths=paths,
            now=datetime(2026, 8, 10, 12, 0, 0),
        )
        discarded_pdf = paths.app_data_dir / discarded["archived_pdf_path"]
        child_batch = import_pdf_batch(
            [
                PdfBatchSource(
                    path=discarded_pdf,
                    original_file_name="documento.pdf",
                    parent_invoice_record_id=discarded["record_id"],
                    source_page_numbers=(1, 3),
                )
            ],
            runtime_paths=paths,
        )

        def child_processor(_source_path: Path, output_dir: Path, first_number: int):
            return _write_manifest(output_dir, first_number, [2])

        child_run = process_imported_batch(
            child_batch.batch_id,
            runtime_paths=paths,
            document_processor=child_processor,
        )
        child_invoice = _invoice_payload("invoice-001")
        child_invoice["pageNumbers"] = [2]
        saved = save_reviewed_invoice(
            {
                "selected_batch_id": child_run.review_batches[0].id,
                "review_batch_base_path": child_run.review_batches[0].base_path,
                "invoice": child_invoice,
            },
            runtime_paths=paths,
            now=datetime(2026, 8, 10, 13, 0, 0),
        )

        assert discarded["display_name"] == "documento - pagine 1, 3"
        assert saved["archived_pdf_page_numbers"] == [3]
        assert len(PdfReader(str(paths.app_data_dir / saved["archived_pdf_path"])).pages) == 1
        assert len(list(paths.archive_sources_dir.glob("*.pdf"))) == 1
    finally:
        _clean_workspace(workspace)


def test_save_invoice_registry_excel_exports_only_saved_records(monkeypatch) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        _mock_pdf_archive(monkeypatch, paths)
        save_reviewed_invoice(
            {
                "selected_batch_id": review_batch.id,
                "review_batch_base_path": review_batch.base_path,
                "invoice": _invoice_payload("invoice-001"),
            },
            runtime_paths=paths,
            now=datetime(2026, 8, 10, 12, 0, 0),
        )
        registry = load_invoice_registry(paths)
        registry["records"].append(
            {
                "record_id": "discarded-record",
                "status": "discarded",
                "invoice": _invoice_payload("invoice-999"),
            }
        )
        paths.invoice_registry_file.write_text(
            json.dumps(registry, ensure_ascii=False),
            encoding="utf-8",
        )

        output_path = save_invoice_registry_excel(runtime_paths=paths)

        workbook = load_workbook(output_path)
        sheet = workbook["Registro fatture"]
        headers = [cell.value for cell in sheet[1]]
        values = [cell.value for cell in sheet[2]]
        assert headers[:4] == [
            "Salvata il",
            "Nome emittente",
            "C.F. emittente",
            "P.IVA emittente",
        ]
        assert sheet.max_row == 2
        assert values[0] == "2026-08-10"
        assert values[6] == "6254/CV"
        assert values[8] == "85,79"
        assert values[12] == "source.pdf"
        assert values[13] == "archive/invoices/fake_invoice.pdf"
        assert values[14] == "archive/sources/source.pdf"
    finally:
        _clean_workspace(workspace)


def test_invoice_registry_excel_reads_manifest_economic_field_aliases(monkeypatch) -> None:
    workspace = _workspace()
    try:
        paths, review_batch = _processed_review_batch(workspace)
        invoice_payload = _invoice_payload("invoice-001")
        invoice_payload["fields"] = [
            {
                "kind": "economic_summary.totale_fattura.value",
                "label": "Totale fattura",
                "value": "56,21",
                "score": 180,
            },
            {
                "kind": "economic_summary.ritenuta.value",
                "label": "Ritenuta",
                "value": "1,84",
                "score": 220,
            },
            {
                "kind": "economic_summary.totale_netto_fattura.value",
                "label": "Totale netto fattura",
                "value": "54,37",
                "score": 180,
            },
        ]
        _mock_pdf_archive(monkeypatch, paths)
        save_reviewed_invoice(
            {
                "selected_batch_id": review_batch.id,
                "review_batch_base_path": review_batch.base_path,
                "invoice": invoice_payload,
            },
            runtime_paths=paths,
            now=datetime(2026, 8, 10, 12, 0, 0),
        )

        output_path = save_invoice_registry_excel(runtime_paths=paths)

        workbook = load_workbook(output_path)
        sheet = workbook["Registro fatture"]
        values = [cell.value for cell in sheet[2]]
        assert values[8] == "56,21"
        assert values[9] == "1,84"
        assert values[10] == "54,37"
    finally:
        _clean_workspace(workspace)


def test_delete_invoice_pdf_preserves_registry_and_source_archive() -> None:
    workspace = _workspace()
    try:
        paths = get_runtime_paths(workspace / "app-data")
        paths.ensure_directories()
        invoice_path = paths.archive_invoices_dir / "invoice.pdf"
        source_path = paths.archive_sources_dir / "source.pdf"
        invoice_path.write_bytes(b"invoice")
        source_path.write_bytes(b"source")
        paths.invoice_registry_file.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "records": [
                        {
                            "record_id": "record-001",
                            "status": "saved",
                            "archived_pdf_path": "archive/invoices/invoice.pdf",
                            "archived_source_pdf_path": "archive/sources/source.pdf",
                            "invoice": {"fields": []},
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )

        result = delete_invoice_archive(
            "invoice",
            record_id="record-001",
            runtime_paths=paths,
        )

        assert result["deleted_invoice_count"] == 1
        assert result["deleted_invoice_file_count"] == 1
        assert result["deleted_source_file_count"] == 0
        assert not invoice_path.exists()
        assert source_path.exists()
        record = load_invoice_registry(paths)["records"][0]
        assert record["invoice_archive_visible"] is False
        assert record["archived_pdf_path"] == ""
        assert record["archived_source_pdf_path"] == "archive/sources/source.pdf"
    finally:
        _clean_workspace(workspace)


def test_delete_registry_row_preserves_both_pdf_archives() -> None:
    workspace = _workspace()
    try:
        paths = get_runtime_paths(workspace / "app-data")
        paths.ensure_directories()
        invoice_path = paths.archive_invoices_dir / "invoice.pdf"
        source_path = paths.archive_sources_dir / "source.pdf"
        invoice_path.write_bytes(b"invoice")
        source_path.write_bytes(b"source")
        paths.invoice_registry_file.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "records": [
                        {
                            "record_id": "record-001",
                            "status": "saved",
                            "archived_pdf_path": "archive/invoices/invoice.pdf",
                            "archived_pdf_url": "/runtime/archive/invoices/invoice.pdf",
                            "archived_source_pdf_path": "archive/sources/source.pdf",
                            "archived_source_pdf_url": "/runtime/archive/sources/source.pdf",
                            "invoice": {"fields": []},
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )

        result = delete_invoice_archive(
            "registry",
            record_ids=["record-001"],
            runtime_paths=paths,
        )

        assert result["deleted_registry_count"] == 1
        assert result["deleted_invoice_file_count"] == 0
        assert result["deleted_source_file_count"] == 0
        assert invoice_path.exists()
        assert source_path.exists()
        record = load_invoice_registry(paths)["records"][0]
        assert record["registry_visible"] is False
        assert record["archived_pdf_path"] == "archive/invoices/invoice.pdf"
        assert record["archived_source_pdf_path"] == "archive/sources/source.pdf"
    finally:
        _clean_workspace(workspace)


def test_purge_discarded_invoice_records_only_removes_legacy_marker_records() -> None:
    workspace = _workspace()
    try:
        paths = get_runtime_paths(workspace / "app-data")
        paths.ensure_directories()
        paths.invoice_registry_file.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "records": [
                        {
                            "record_id": "legacy-marker",
                            "status": "discarded",
                            "archived_pdf_path": "",
                        },
                        {
                            "record_id": "persisted-discard",
                            "status": "discarded",
                            "archived_pdf_path": "archive/invoices/scartata.pdf",
                        },
                        {
                            "record_id": "saved-record",
                            "status": "saved",
                            "archived_pdf_path": "archive/invoices/salvata.pdf",
                        },
                    ],
                }
            ),
            encoding="utf-8",
        )

        assert purge_discarded_invoice_records(paths) == 1

        registry = load_invoice_registry(paths)
        assert registry["schema_version"] == 3
        assert [record["record_id"] for record in registry["records"]] == [
            "persisted-discard",
            "saved-record",
        ]
        assert purge_discarded_invoice_records(paths) == 0
    finally:
        _clean_workspace(workspace)


def test_load_invoice_registry_migrates_legacy_withholding_percentage_to_amount() -> None:
    workspace = _workspace()
    try:
        paths = get_runtime_paths(workspace / "app-data")
        paths.ensure_directories()
        paths.invoice_registry_file.write_text(
            json.dumps(
                {
                    "schema_version": 2,
                    "records": [
                        {
                            "record_id": "legacy-withholding",
                            "invoice": {
                                "fields": [
                                    {"kind": "totale_fattura", "value": "100,00"},
                                    {
                                        "kind": "ritenuta",
                                        "label": "Ritenuta %",
                                        "value": "4%",
                                    },
                                    {"kind": "totale_netto_fattura", "value": "96,00"},
                                ]
                            },
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )

        registry = load_invoice_registry(paths)
        fields = registry["records"][0]["invoice"]["fields"]
        withholding = next(field for field in fields if field["kind"] == "ritenuta")

        assert registry["schema_version"] == 3
        assert withholding["label"] == "Ritenuta"
        assert withholding["value"] == "4.00"
    finally:
        _clean_workspace(workspace)


def _processed_review_batch(workspace: Path):
    source_file = workspace / "source.pdf"
    source_file.write_bytes(b"%PDF")
    paths = get_runtime_paths(workspace / "app-data")
    imported_batch = import_pdf_batch(
        [source_file],
        runtime_paths=paths,
        now=datetime(2026, 8, 10, 10, 0, 0),
    )

    def fake_processor(_source_path: Path, output_dir: Path, first_invoice_number: int):
        output_dir.mkdir(parents=True, exist_ok=True)
        manifest_path = output_dir / "viewer_manifest.json"
        manifest_path.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "source_file": "source.pdf",
                    "pages": [],
                    "invoices": [
                        {
                            "invoice_id": f"invoice-{first_invoice_number:03d}",
                            "title": f"Fattura {first_invoice_number:03d}",
                            "page_numbers": [1],
                            "fields": [],
                        }
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        return [manifest_path]

    run_index = process_imported_batch(
        imported_batch.batch_id,
        runtime_paths=paths,
        document_processor=fake_processor,
        now=datetime(2026, 8, 10, 11, 0, 0),
    )
    return paths, run_index.review_batches[0]


def _invoice_payload(invoice_id: str) -> dict:
    return {
        "invoiceId": invoice_id,
        "title": "Fattura 001",
        "documentId": "doc_000001",
        "manifestInvoiceId": invoice_id.rsplit(":", 1)[-1],
        "sourceFile": "source.pdf",
        "basePath": "/runtime/manifests/process_20260810_110000/documents/doc_000001",
        "pageNumbers": [1],
        "pages": [],
        "fields": [
            {
                "kind": "totale_fattura",
                "label": "Totale fattura",
                "value": "85,79",
                "score": 0.96,
            },
            {
                "kind": "numero_fattura",
                "label": "Numero fattura",
                "value": "6254/CV",
                "score": 0.91,
            }
        ],
    }


def _record_with_field_value(record: dict, field_kind: str, value: str) -> dict:
    return {
        **record,
        "invoice": {
            **record["invoice"],
            "fields": [
                {
                    **field,
                    "value": value if field["kind"] == field_kind else field["value"],
                }
                for field in record["invoice"]["fields"]
            ],
        },
    }


def _invoice_field_values_for_test(record: dict) -> dict[str, str]:
    return {
        str(field["kind"]): str(field.get("value") or "")
        for field in record["invoice"]["fields"]
    }


def _mock_pdf_archive(monkeypatch, paths) -> None:
    def fake_archive(
        _invoice_payload,
        _invoice_ref,
        _runtime_paths,
        **_kwargs,
    ):
        return ArchivedInvoicePdf(
            path=paths.archive_invoices_dir / "fake_invoice.pdf",
            relative_path="archive/invoices/fake_invoice.pdf",
            page_numbers=[1],
            source_pdf_path=paths.batches_dir / "batch" / "input" / "doc_000001.pdf",
            archived_source_path=paths.archive_sources_dir / "source.pdf",
            archived_source_relative_path="archive/sources/source.pdf",
        )

    monkeypatch.setattr(
        "estrattore.invoice_registry_storage.archive_invoice_pdf",
        fake_archive,
    )


def _write_blank_pdf(path: Path, page_count: int) -> None:
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=72, height=72)
    with path.open("wb") as output_file:
        writer.write(output_file)


def _write_manifest(output_dir: Path, first_number: int, page_numbers: list[int]):
    output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = output_dir / "viewer_manifest.json"
    manifest_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "source_file": "documento.pdf",
                "pages": [],
                "invoices": [
                    {
                        "invoice_id": f"invoice-{first_number:03d}",
                        "title": f"Fattura {first_number:03d}",
                        "page_numbers": page_numbers,
                        "fields": [],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    return [manifest_path]
